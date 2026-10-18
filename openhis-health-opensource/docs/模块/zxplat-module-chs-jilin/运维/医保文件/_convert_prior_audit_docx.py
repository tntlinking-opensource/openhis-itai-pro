# -*- coding: utf-8 -*-
"""Convert 全省事前联调方案 docx → markdown."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

SRC = Path(__file__).with_name("附件：全省事前联调方案1.5.docx")
OUT = Path(__file__).with_name("附件：全省事前联调方案1.5.md")


def cell_text(cell) -> str:
    texts = []
    for p in cell.paragraphs:
        t = (p.text or "").strip()
        if t:
            texts.append(t)
    return "<br>".join(texts).replace("|", "\\|")


def table_to_md(table) -> str:
    rows = []
    for row in table.rows:
        rows.append([cell_text(c) for c in row.cells])
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    norm = [r + [""] * (cols - len(r)) for r in rows]
    lines = [
        "| " + " | ".join(norm[0]) + " |",
        "| " + " | ".join(["---"] * cols) + " |",
    ]
    for r in norm[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def style_name(p) -> str:
    try:
        return p.style.name if p.style else ""
    except Exception:
        return ""


def heading_level(style: str, text: str) -> int | None:
    sn = (style or "").lower()
    m = re.search(r"heading\s*(\d+)", sn)
    if m:
        return int(m.group(1))
    # Word 中文样式常见名
    for i in range(1, 6):
        if f"标题 {i}" in style or f"标题{i}" in style:
            return i
    t = text.strip()
    if re.match(r"^第[一二三四五六七八九十百零0-9]+[章节部分]", t):
        return 1
    if re.match(r"^[一二三四五六七八九十]+[、.．]", t):
        return 2
    if re.match(r"^（[一二三四五六七八九十]+）", t) or re.match(r"^\([一二三四五六七八九十]+\)", t):
        return 3
    m2 = re.match(r"^(\d+(?:\.\d+)*)\s+\S", t)
    if m2 and len(t) < 100:
        return min(m2.group(1).count(".") + 2, 5)
    return None


def main() -> None:
    doc = Document(str(SRC))
    body_elements: list[tuple[str, object]] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            for p in doc.paragraphs:
                if p._element is child:
                    body_elements.append(("p", p))
                    break
        elif tag == "tbl":
            for t in doc.tables:
                if t._element is child:
                    body_elements.append(("t", t))
                    break

    if not body_elements:
        for p in doc.paragraphs:
            body_elements.append(("p", p))
        for t in doc.tables:
            body_elements.append(("t", t))

    parts: list[str] = [
        "---",
        'source: "附件：全省事前联调方案1.5.docx"',
        'converted: "2026-07-23"',
        'note: "由 python-docx 自动转换，表格合并单元格可能重复，需人工校对"',
        "---",
        "",
        "# 全省事前联调方案 1.5",
        "",
    ]

    for kind, obj in body_elements:
        if kind == "p":
            text = (obj.text or "").strip()
            if not text:
                continue
            lvl = heading_level(style_name(obj), text)
            if lvl:
                parts.append("#" * min(lvl + 1, 6) + " " + text)
            else:
                parts.append(text)
            parts.append("")
        else:
            md = table_to_md(obj)
            if md:
                parts.append(md)
                parts.append("")

    OUT.write_text("\n".join(parts), encoding="utf-8")
    print(f"wrote {OUT}")
    print(f"size={OUT.stat().st_size} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
