# -*- coding: utf-8 -*-
import collections
from openpyxl import load_workbook
import psycopg2

p = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\参考文件\IDC-10临床2.0版与国家医保版本映射.xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
ws = wb["对照库"]
rows = ws.iter_rows(values_only=True)
next(rows)
by_clinic = collections.defaultdict(set)
empty = 0
for clinic, cname, yb, yname in rows:
    clinic = str(clinic or "").strip()
    yb = str(yb or "").strip()
    if not clinic or not yb:
        empty += 1
        continue
    by_clinic[clinic].add(yb)
wb.close()
unique = {c: next(iter(ys)) for c, ys in by_clinic.items() if len(ys) == 1}
multi = {c: ys for c, ys in by_clinic.items() if len(ys) > 1}
print("rows_used", sum(len(v) for v in by_clinic.values()), "clinics", len(by_clinic), "unique", len(unique), "multi", len(multi), "blank", empty)
if multi:
    sample = list(multi.items())[:5]
    print("multi sample", [(k, list(v)[:4]) for k, v in sample])

conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname="zxplat-bak")
cur = conn.cursor()
cur.execute("""
    SELECT condition_code, coalesce(yb_no,'')
    FROM cln_condition_definition
    WHERE is_deleted = false AND source_enum = 'clinical_v2'
""")
cv_empty = cv_fill = cv_same = cv_diff = 0
for code, yb in cur.fetchall():
    code = (code or "").strip()
    yb = yb.strip()
    mapped = unique.get(code)
    if not yb:
        cv_empty += 1
        if mapped:
            cv_fill += 1
    elif mapped:
        if mapped == yb:
            cv_same += 1
        else:
            cv_diff += 1
print("cv2 empty", cv_empty, "empty_can_fill", cv_fill, "existing_same", cv_same, "existing_diff", cv_diff)

cur.execute("""
    SELECT condition_code, coalesce(clinic_code,''), coalesce(yb_no,'')
    FROM cln_condition_definition
    WHERE is_deleted = false AND source_enum = 'general_dx'
""")
g_no_clinic = g_clinic_fill = g_yb_upgrade = g_yb_conflict = 0
for code, clinic, yb in cur.fetchall():
    code = (code or "").strip()
    clinic = clinic.strip()
    yb = yb.strip()
    mapped = unique.get(code) or unique.get(clinic)
    if not clinic and code in unique:
        g_clinic_fill += 1
    if not clinic:
        g_no_clinic += 1
    if mapped and (not yb or yb == code) and mapped != yb:
        g_yb_upgrade += 1
    if mapped and yb and yb != code and yb != mapped:
        g_yb_conflict += 1
print("gdx no clinic", g_no_clinic, "can clinic", g_clinic_fill, "yb upgrade", g_yb_upgrade, "yb conflict", g_yb_conflict)
conn.close()
