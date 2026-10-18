# -*- coding: utf-8 -*-
"""用临床2.0与国家医保版映射，补齐并校正诊断对照。"""
import collections

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import execute_values

XLSX = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\参考文件\IDC-10临床2.0版与国家医保版本映射.xlsx"
DBS = ("zxplat-bak", "whale-training")


def load_map():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    rows = wb["对照库"].iter_rows(values_only=True)
    next(rows)
    grouped = collections.defaultdict(set)
    for clinic, _cname, yb, _yname in rows:
        clinic = str(clinic or "").strip()
        yb = str(yb or "").strip()
        if clinic and yb:
            grouped[clinic].add(yb)
    wb.close()
    unique = [(clinic, next(iter(codes))) for clinic, codes in grouped.items() if len(codes) == 1]
    print("unique", len(unique), "ambiguous", sum(1 for codes in grouped.values() if len(codes) > 1))
    return unique


def apply_db(db, pairs):
    conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname=db)
    conn.autocommit = False
    cur = conn.cursor()
    cur.execute("SELECT current_database()")
    if cur.fetchone()[0] != db:
        raise SystemExit("库名不符: " + db)
    cur.execute("""
        CREATE TEMP TABLE tmp_icd_map (
            clinic_code varchar(64) PRIMARY KEY,
            yb_no varchar(64) NOT NULL
        ) ON COMMIT DROP
    """)
    execute_values(cur, "INSERT INTO tmp_icd_map (clinic_code, yb_no) VALUES %s", pairs)

    cur.execute("""
        UPDATE cln_condition_definition d
        SET yb_no = m.yb_no,
            yb_flag = TRUE,
            yb_match_flag = TRUE,
            updated_at = CURRENT_TIMESTAMP,
            updated_by = 1
        FROM tmp_icd_map m
        WHERE d.is_deleted = FALSE
          AND d.source_enum = 'clinical_v2'
          AND d.condition_code = m.clinic_code
          AND btrim(coalesce(d.yb_no, '')) IS DISTINCT FROM m.yb_no
    """)
    cv2 = cur.rowcount

    cur.execute("""
        UPDATE cln_condition_definition d
        SET clinic_code = CASE
                WHEN d.clinic_code IS NULL OR btrim(d.clinic_code) = '' THEN d.condition_code
                ELSE d.clinic_code
            END,
            yb_no = m.yb_no,
            yb_flag = TRUE,
            yb_match_flag = TRUE,
            updated_at = CURRENT_TIMESTAMP,
            updated_by = 1
        FROM tmp_icd_map m
        WHERE d.is_deleted = FALSE
          AND d.source_enum <> 'clinical_v2'
          AND m.clinic_code = COALESCE(NULLIF(btrim(d.clinic_code), ''), btrim(d.condition_code))
          AND (
                d.clinic_code IS NULL OR btrim(d.clinic_code) = ''
                OR d.yb_no IS NULL OR btrim(d.yb_no) = ''
                OR btrim(d.yb_no) IS DISTINCT FROM m.yb_no
          )
    """)
    other = cur.rowcount
    conn.commit()

    cur.execute("""
        SELECT source_enum,
               COUNT(*),
               SUM(CASE WHEN clinic_code IS NULL OR btrim(clinic_code) = '' THEN 1 ELSE 0 END),
               SUM(CASE WHEN yb_no IS NULL OR btrim(yb_no) = '' THEN 1 ELSE 0 END)
        FROM cln_condition_definition
        WHERE is_deleted = FALSE AND source_enum IN ('clinical_v2', 'general_dx')
        GROUP BY source_enum
        ORDER BY source_enum
    """)
    print(db, "cv2_updated", cv2, "other_updated", other)
    for row in cur.fetchall():
        print(" ", row)
    conn.close()


def main():
    pairs = load_map()
    for db in DBS:
        apply_db(db, pairs)


if __name__ == "__main__":
    main()
