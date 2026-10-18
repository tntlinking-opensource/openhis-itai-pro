# -*- coding: utf-8 -*-
"""从工作目录导出当前启用的服务项目清单。"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(__file__).with_name("service_catalog_working.csv")
OUT = Path(__file__).with_name("当前可用服务项目.xlsx")

LV = {"1": "甲", "2": "乙", "3": "丙"}
HEADERS = [
    "口径", "项目代码", "医保目录编码", "项目名称", "收费等级", "计价单位", "定价上限",
    "地方编码", "限制使用标志", "限制使用范围", "诊疗项目说明", "诊疗除外内容",
    "诊疗项目内涵", "开始日期",
]
SRC_COLS = [
    "catalog_kind", "project_code", "hilist_code", "hilist_name", "chrgitm_lv",
    "计价单位", "定价上限", "地方编码", "限制使用标志", "限制使用范围",
    "诊疗项目说明", "诊疗除外内容", "诊疗项目内涵", "begin_date",
]


def main():
    with SRC.open("r", encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.DictReader(f) if (r.get("status") or "").strip() == "ACTIVE"]
    rows.sort(key=lambda r: (
        (r.get("project_code") or "").strip(),
        (r.get("hilist_code") or "").strip(),
        r.get("catalog_kind") or "",
    ))
    n_f = sum(1 for r in rows if r.get("catalog_kind") == "SERVICE_F")
    n_s = sum(1 for r in rows if r.get("catalog_kind") == "SERVICE")

    wb = Workbook()
    ws = wb.active
    ws.title = "可用项目"
    ws.append(HEADERS)
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True)
    for rec in rows:
        vals = []
        for col in SRC_COLS:
            value = (rec.get(col) or "").strip()
            if col == "chrgitm_lv":
                value = LV.get(value, value)
            if col == "catalog_kind":
                value = "F服务" if value == "SERVICE_F" else "普通服务"
            vals.append(value)
        ws.append(vals)
    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    ws.freeze_panes = "A2"
    widths = [12, 22, 36, 36, 10, 14, 12, 16, 14, 40, 40, 28, 48, 14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    note = wb.create_sheet("说明")
    note.append(["来源", str(SRC)])
    note.append(["口径", "只保留 status=ACTIVE。STOPPED 未收入。"])
    note.append(["基准日", "2026-10-07。启用行均无已过止期。"])
    note.append(["条数", f"{len(rows)}（F服务 {n_f}，普通服务 {n_s}）"])
    note.append(["重复", "12 个编码在普通服务与 F服务 同时启用，限价不同，两条都保留，用「口径」区分。"])
    note.append(["收费等级", "1甲 2乙 3丙"])
    note.column_dimensions["A"].width = 12
    note.column_dimensions["B"].width = 88
    wb.save(OUT)
    print(f"WROTE {OUT}")
    print(f"rows {len(rows)} F {n_f} S {n_s}")


if __name__ == "__main__":
    main()
