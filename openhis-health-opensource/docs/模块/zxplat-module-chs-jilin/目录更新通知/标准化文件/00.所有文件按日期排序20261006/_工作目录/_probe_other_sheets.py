# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import psycopg2

XLSX = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\数据迁移相关\服务目录\国基医院新系统诊疗目录表.xlsx"
wb = load_workbook(XLSX, read_only=True, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    print(name, "rows", len(rows), "head", rows[0] if rows else None)
    if len(rows) > 1:
        print("  row2", rows[1])
wb.close()

conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname="whale-training")
cur = conn.cursor()
cur.execute("""
SELECT bus_no, name, coalesce(yb_no,'')
FROM wkf_activity_definition
WHERE is_deleted = false
  AND (mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '')
ORDER BY name
""")
empty = cur.fetchall()
conn.close()
print("empty", len(empty))
# print category-like prefixes
from collections import Counter
c = Counter()
for bus, name, yb in empty:
    if "诊查" in name:
        c["诊查"] += 1
    elif "肝功" in name or "套" in name or "项）" in name or name.endswith("项"):
        c["组合/项"] += 1
    elif "CT" in name or "DR" in name:
        c["影像部位"] += 1
    else:
        c["其他"] += 1
print(c)
print("--- other names ---")
n = 0
for bus, name, yb in empty:
    if "诊查" in name or "肝功" in name or "CT" in name or "DR" in name:
        continue
    print(bus, name, yb)
    n += 1
    if n >= 40:
        break
