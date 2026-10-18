# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import psycopg2

XLSX = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\数据迁移相关\服务目录\国基医院新系统诊疗目录表.xlsx"
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["②诊疗目录"]
rows = ws.iter_rows(values_only=True)
next(rows)
header = list(next(rows))
name_i = header.index("*诊疗名称")
yb_i = header.index("医保编号")
mr_i = header.index("病案结算类型")
file_names = set()
file_yb = []
want = ("门诊诊查费", "肝功七项", "腰椎正位DR", "氩气刀")
hits = []
for row in rows:
    name = str(row[name_i] or "").strip()
    if not name or name.startswith("[示例]"):
        continue
    file_names.add(name)
    yb = str(row[yb_i] or "").strip()
    if yb:
        file_yb.append(yb)
    if any(key in name for key in want) and len(hits) < 20:
        hits.append((name, yb, row[mr_i]))
print("file names", len(file_names), "with yb", len(file_yb))
print("file hits", hits)
print("yb samples", file_yb[:8])
wb.close()

conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname="whale-training")
cur = conn.cursor()
cur.execute("""
SELECT name, yb_no FROM wkf_activity_definition
WHERE is_deleted = false AND (mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '')
  AND (name LIKE '%诊查费%' OR name LIKE '%肝功%' OR name LIKE '%氩气%' OR yb_no LIKE '011102020010000%')
LIMIT 20
""")
print("--- empty ad ---")
for row in cur.fetchall():
    print(row)
cur.execute("""
SELECT
  SUM(CASE WHEN yb_no IS NULL OR btrim(yb_no) = '' THEN 1 ELSE 0 END),
  SUM(CASE WHEN yb_no IS NOT NULL AND btrim(yb_no) <> '' THEN 1 ELSE 0 END)
FROM wkf_activity_definition
WHERE is_deleted = false AND (mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '')
""")
print("empty mr no-yb / has-yb", cur.fetchone())
cur.execute("""
SELECT left(yb_no, 20), COUNT(*) FROM wkf_activity_definition
WHERE is_deleted = false AND yb_no IS NOT NULL AND btrim(yb_no) <> ''
GROUP BY 1 ORDER BY COUNT(*) DESC LIMIT 5
""")
print("ad yb prefix", cur.fetchall())
conn.close()

prefixes = set()
for yb in file_yb:
    prefixes.add(yb.split("-")[0])
print("file distinct code-before-dash", len(prefixes))
print("example prefixes", list(prefixes)[:5])
