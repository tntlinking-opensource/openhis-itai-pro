# -*- coding: utf-8 -*-
from collections import Counter
from openpyxl import load_workbook
import psycopg2

p = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\数据迁移相关\服务目录\国基医院新系统诊疗目录表.xlsx"
wb = load_workbook(p, read_only=True, data_only=True)
ws = wb["②诊疗目录"]
rows = ws.iter_rows(values_only=True)
next(rows)
header = list(next(rows))
print("header", header)
fin_i = header.index("*财务类型")
mr_i = header.index("病案结算类型")
name_i = header.index("*诊疗名称")
yb_i = header.index("医保编号")
legacy_i = header.index("院内旧编码")
cat_i = header.index("*目录分类")
fin_c = Counter()
mr_c = Counter()
both_empty = fin_filled = mr_filled = data = 0
samples = []
for row in rows:
    name = str(row[name_i] or "").strip()
    if not name or name.startswith("[示例]"):
        continue
    data += 1
    fin = str(row[fin_i] or "").strip()
    mr = str(row[mr_i] or "").strip()
    fin_c[fin or "(空)"] += 1
    mr_c[mr or "(空)"] += 1
    if fin:
        fin_filled += 1
    if mr:
        mr_filled += 1
    if not fin and not mr:
        both_empty += 1
    if len(samples) < 3 and fin and mr:
        samples.append((name, row[cat_i], row[yb_i], row[legacy_i], fin, mr))
print("data", data, "fin_filled", fin_filled, "mr_filled", mr_filled, "both_empty", both_empty)
print("fin", fin_c.most_common(20))
print("mr", mr_c.most_common(20))
print("samples", samples)
for sheet in ["旧系统手术", "旧系统护理", "旧系统治疗费", "2026年5月25启用项目"]:
    ws2 = wb[sheet]
    it = ws2.iter_rows(values_only=True)
    h = list(next(it))
    print("SHEET", sheet, h)
wb.close()

conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname="whale-training")
cur = conn.cursor()
cur.execute("SELECT COUNT(*) FROM wkf_activity_definition WHERE is_deleted = false")
print("ad", cur.fetchone()[0])
cur.execute("""
SELECT
  SUM(CASE WHEN mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '' THEN 1 ELSE 0 END),
  SUM(CASE WHEN mr_settlement_type_code IS NOT NULL AND btrim(mr_settlement_type_code) <> '' THEN 1 ELSE 0 END)
FROM wkf_activity_definition WHERE is_deleted = false
""")
print("mr empty/filled", cur.fetchone())
cur.execute("""
SELECT column_name, is_nullable, character_maximum_length
FROM information_schema.columns
WHERE table_name = 'adm_charge_item_definition' AND column_name IN ('type_code','instance_table','instance_id','price_type_code','is_deleted')
""")
print("cid cols", cur.fetchall())
cur.execute("""
SELECT
  COUNT(*) AS price_rows,
  SUM(CASE WHEN d.type_code IS NULL OR btrim(d.type_code) = '' THEN 1 ELSE 0 END) AS empty_type
FROM adm_charge_item_definition d
JOIN wkf_activity_definition a ON a.id = d.instance_id AND a.is_deleted = false
WHERE d.is_deleted = false AND d.instance_table = 'wkf_activity_definition'
""")
print("price empty type", cur.fetchone())
cur.execute("""
SELECT d.price_type_code,
       SUM(CASE WHEN d.type_code IS NULL OR btrim(d.type_code) = '' THEN 1 ELSE 0 END),
       COUNT(*)
FROM adm_charge_item_definition d
JOIN wkf_activity_definition a ON a.id = d.instance_id AND a.is_deleted = false
WHERE d.is_deleted = false AND d.instance_table = 'wkf_activity_definition'
GROUP BY d.price_type_code
""")
print("by price type", cur.fetchall())
cur.execute("""
SELECT COUNT(*) FROM wkf_activity_definition a
WHERE a.is_deleted = false
  AND NOT EXISTS (
    SELECT 1 FROM adm_charge_item_definition d
    WHERE d.instance_table = 'wkf_activity_definition' AND d.instance_id = a.id AND d.is_deleted = false
  )
""")
print("ad without price", cur.fetchone()[0])
cur.execute("""
SELECT cs.name, c.code, c.display
FROM ter_concept c
JOIN ter_code_system cs ON cs.id = c.code_system_id
WHERE cs.name IN ('fin_type_code', 'mr_settlement_type_code')
  AND c.is_deleted = false AND cs.is_deleted = false
ORDER BY cs.name, c.sort_order NULLS LAST, c.code
""")
print("--- terms ---")
for row in cur.fetchall():
    print(row)
cur.execute("""
SELECT a.bus_no, a.name, a.yb_no, left(coalesce(a.description_text,''), 80), a.mr_settlement_type_code
FROM wkf_activity_definition a
WHERE a.is_deleted = false
LIMIT 3
""")
print("ad sample", cur.fetchall())
conn.close()
