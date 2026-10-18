# -*- coding: utf-8 -*-
"""whale-training 诊疗 AD 对照 CSV（修改前快照）。只读库，不改数据。"""
from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CSV_PATH = Path(
    r"D:\0.product\03.ZXPlat\zxplat\docs\模块\zxplat-module-chs-jilin"
    r"\目录更新通知\标准化文件\00.所有文件按日期排序20261006\_工作目录"
    r"\service_catalog_working.csv"
)
OUT_PATH = CSV_PATH.with_name("whale-training-AD对照-修改前.xlsx")

DB = dict(
    host="192.168.1.1",
    port=5433,
    dbname="whale-training",
    user="postgres",
    password="root",
)

SKIP_KEYS = {"国家未赋码"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
FAIL_FILL = PatternFill("solid", fgColor="F8CBAD")


def blank(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() == "null" or s == "--":
        return None
    return s


def norm_name(v):
    s = blank(v)
    if s is None:
        return None
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\u3000", " ")
    s = re.sub(r"\s+", "", s)
    return s


def catalog_key(row):
    return blank(row.get("project_code")) or blank(row.get("hilist_code"))


def row_codes(row):
    codes = []
    for col in ("project_code", "hilist_code", "地方编码"):
        c = blank(row.get(col))
        if c and c not in SKIP_KEYS and c not in codes:
            codes.append(c)
    return codes


def parse_lv(v):
    s = blank(v)
    if s is None:
        return None
    try:
        return int(Decimal(s))
    except (InvalidOperation, ValueError):
        return None


def parse_amt(v):
    s = blank(v)
    if s is None:
        return None
    try:
        n = Decimal(s)
    except InvalidOperation:
        return None
    if n <= 0:
        return None
    return n


def load_csv():
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fetch_ads(cur):
    cur.execute(
        """
        SELECT id, tenant_id, bus_no, name, merchandise_name, yb_no, yb_flag, yb_match_flag,
               chrgitm_lv, status_enum, package_unit_code, description_text, is_deleted
          FROM wkf_activity_definition
         WHERE coalesce(is_deleted, false) = false
        """
    )
    cols = [d[0] for d in cur.description]
    rows = []
    by_code = defaultdict(list)
    for tup in cur.fetchall():
        rec = dict(zip(cols, tup))
        rec["id"] = rec["id"]
        rows.append(rec)
        for key in (blank(rec.get("yb_no")), blank(rec.get("bus_no"))):
            if key:
                by_code[key].append(rec)
    return rows, by_code


def fetch_sale_prices(cur):
    cur.execute(
        """
        SELECT instance_id, min(price_value) AS min_sale, max(price_value) AS max_sale
          FROM adm_charge_item_definition
         WHERE coalesce(is_deleted, false) = false
           AND instance_table = 'wkf_activity_definition'
           AND (status_enum IS NULL OR status_enum = 'ACTIVE')
           AND price_type_code = 'SALE'
         GROUP BY instance_id
        """
    )
    out = {}
    for instance_id, min_sale, max_sale in cur.fetchall():
        out[instance_id] = (min_sale, max_sale)
    return out


def fetch_overlay(cur):
    cur.execute(
        """
        SELECT hilist_code, list_type, notice_avali_flag, notice_end_date
          FROM yb_national_catalog_notice_overlay
         WHERE coalesce(is_deleted, false) = false
        """
    )
    by_code = defaultdict(list)
    for hilist_code, list_type, flag, end_date in cur.fetchall():
        code = blank(hilist_code)
        if code:
            by_code[code].append(
                {
                    "list_type": list_type,
                    "notice_avali_flag": blank(flag),
                    "notice_end_date": end_date,
                }
            )
    return by_code


def fetch_unit_labels(cur, codes):
    codes = [c for c in codes if c]
    if not codes:
        return {}
    try:
        cur.execute(
            """
            SELECT bus_no, display
              FROM adm_concept
             WHERE coalesce(is_deleted, false) = false
               AND bus_no = ANY(%s)
            """,
            (codes,),
        )
        return {blank(b): blank(d) for b, d in cur.fetchall() if blank(b)}
    except Exception:
        cur.connection.rollback()
        return {}


def match_ads(codes, by_code):
    seen = set()
    hits = []
    for c in codes:
        for ad in by_code.get(c, []):
            if ad["id"] in seen:
                continue
            seen.add(ad["id"])
            hits.append(ad)
    return hits


def overlay_for(codes, by_ov):
    for c in codes:
        if c in by_ov:
            return by_ov[c]
    return []


def write_sheet(wb, title, headers, rows, fail_col=None):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True)
    for row in rows:
        ws.append(list(row))
        if fail_col is not None:
            val = row[fail_col] if fail_col < len(row) else None
            fill = None
            if val in ("失败", "未对码", "超上限", "仍启用"):
                fill = FAIL_FILL
            elif val in ("警告", "名称不等", "一对多"):
                fill = WARN_FILL
            if fill:
                for cell in ws[ws.max_row]:
                    cell.fill = fill
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(28, max(12, len(str(h)) + 4))
    return ws


def main():
    csv_rows = load_csv()
    conn = psycopg2.connect(**DB)
    cur = conn.cursor()
    ads, by_code = fetch_ads(cur)
    prices = fetch_sale_prices(cur)
    overlays = fetch_overlay(cur)
    unit_codes = {blank(a.get("package_unit_code")) for a in ads}
    unit_labels = fetch_unit_labels(cur, list(unit_codes))
    conn.close()

    detail = []
    unmatched_active = []
    stopped_still_on = []
    lv_mismatch = []
    name_drift = []
    over_cap = []
    multi_hit = []

    csv_all_codes = set()
    for row in csv_rows:
        k = catalog_key(row)
        if k and k not in SKIP_KEYS:
            csv_all_codes.add(k)
        for c in row_codes(row):
            csv_all_codes.add(c)

    for i, row in enumerate(csv_rows, 2):
        kind = blank(row.get("catalog_kind"))
        st = blank(row.get("status")) or ""
        project = blank(row.get("project_code"))
        hilist = blank(row.get("hilist_code"))
        name = blank(row.get("hilist_name"))
        lv = parse_lv(row.get("chrgitm_lv"))
        cap = parse_amt(row.get("定价上限"))
        unit = blank(row.get("计价单位"))
        key = catalog_key(row)
        codes = row_codes(row)
        if project in SKIP_KEYS:
            detail.append(
                [
                    i,
                    kind,
                    st,
                    key,
                    project,
                    hilist,
                    name,
                    lv,
                    unit,
                    cap,
                    "跳过",
                    "国家未赋码",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
            continue
        hits = match_ads(codes, by_code)
        ovs = overlay_for(codes, overlays)
        ov_flag = ovs[0]["notice_avali_flag"] if ovs else None
        ov_end = ovs[0]["notice_end_date"] if ovs else None
        issues = []
        result = "通过"
        if not hits:
            if st == "ACTIVE":
                result = "未对码"
                issues.append("ACTIVE无AD")
                unmatched_active.append(
                    [i, kind, key, project, hilist, name, lv, unit, cap, blank(row.get("地方编码"))]
                )
            else:
                result = "通过"
                issues.append("STOPPED无AD")
            ad_ids = None
            ad_bus = None
            ad_name = None
            ad_yb = None
            ad_st = None
            ad_lv = None
            ad_unit = None
            ad_unit_disp = None
            min_sale = None
            max_sale = None
        else:
            if len(hits) > 1:
                issues.append("一对多")
                result = "警告"
                multi_hit.append(
                    [
                        i,
                        kind,
                        key,
                        name,
                        ",".join(str(h["id"]) for h in hits),
                        ",".join(blank(h.get("bus_no")) or "" for h in hits),
                    ]
                )
            names_ok = False
            lv_fail = False
            status_issue = False
            cap_fail = False
            ad_ids = ",".join(str(h["id"]) for h in hits)
            ad_bus = ",".join(blank(h.get("bus_no")) or "" for h in hits)
            ad_name = " | ".join(blank(h.get("name")) or "" for h in hits)
            ad_yb = ",".join(blank(h.get("yb_no")) or "" for h in hits)
            ad_st = ",".join(blank(h.get("status_enum")) or "" for h in hits)
            ad_lv = ",".join("" if h.get("chrgitm_lv") is None else str(h.get("chrgitm_lv")) for h in hits)
            ad_unit = ",".join(blank(h.get("package_unit_code")) or "" for h in hits)
            ad_unit_disp = ",".join(
                unit_labels.get(blank(h.get("package_unit_code")), blank(h.get("package_unit_code")) or "")
                for h in hits
            )
            sales = []
            for h in hits:
                p = prices.get(h["id"])
                if p:
                    sales.append(p)
            min_sale = min((p[0] for p in sales), default=None)
            max_sale = max((p[1] for p in sales), default=None)
            csv_nn = norm_name(name)
            for h in hits:
                if csv_nn and (
                    csv_nn == norm_name(h.get("name")) or csv_nn == norm_name(h.get("merchandise_name"))
                ):
                    names_ok = True
                    break
            if csv_nn and not names_ok:
                issues.append("名称不等")
                if result == "通过":
                    result = "警告"
                name_drift.append([i, kind, key, name, ad_name, ad_ids, ad_bus])
            if lv is not None:
                hit_lvs = {h.get("chrgitm_lv") for h in hits}
                if lv not in hit_lvs:
                    lv_fail = True
                    issues.append("等级不等")
                    result = "失败"
                    lv_mismatch.append([i, kind, key, name, lv, ad_lv, ad_ids, ad_yb])
            if st == "ACTIVE":
                if not any(blank(h.get("status_enum")) == "ACTIVE" for h in hits):
                    status_issue = True
                    issues.append("ACTIVE但AD未启用")
                    result = "失败"
            if st == "STOPPED":
                if any(blank(h.get("status_enum")) == "ACTIVE" for h in hits):
                    status_issue = True
                    issues.append("STOPPED但AD仍启用")
                    if result == "通过":
                        result = "警告"
                    stopped_still_on.append(
                        [i, kind, key, name, ad_ids, ad_bus, ad_yb, ad_st, ov_flag, ov_end]
                    )
            if cap is not None and max_sale is not None and Decimal(str(max_sale)) > cap:
                cap_fail = True
                issues.append("售价超上限")
                result = "失败"
                over_cap.append([i, kind, key, name, cap, max_sale, ad_ids])
            _ = (lv_fail, status_issue, cap_fail)

        detail.append(
            [
                i,
                kind,
                st,
                key,
                project,
                hilist,
                name,
                lv,
                unit,
                cap,
                result,
                ";".join(issues) if issues else "",
                ad_ids,
                ad_bus,
                ad_yb,
                ad_name,
                ad_st,
                ad_lv,
                ad_unit,
                ad_unit_disp,
                min_sale,
                max_sale,
                ov_flag,
                ov_end,
                len(hits) if hits else 0,
            ]
        )

    extra_ad = []
    for ad in ads:
        yb = blank(ad.get("yb_no"))
        bus = blank(ad.get("bus_no"))
        if not yb:
            continue
        if yb in csv_all_codes or (bus and bus in csv_all_codes):
            continue
        extra_ad.append(
            [
                ad["id"],
                bus,
                yb,
                ad.get("name"),
                ad.get("status_enum"),
                ad.get("chrgitm_lv"),
            ]
        )

    counts = defaultdict(int)
    for d in detail:
        counts[d[10]] += 1

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "说明"
    notes = [
        ["库", "whale-training @ 192.168.1.1:5433"],
        ["CSV", str(CSV_PATH)],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["性质", "修改前只读快照，未改库"],
        ["默认口径", "ACTIVE须命中AD；等级须相等；名称不等为警告；STOPPED仍启用记警告（overlay可另停）"],
        ["匹配", "CSV project_code/hilist_code/地方编码 ↔ AD.yb_no 或 AD.bus_no"],
        ["定价", "CID instance=AD 且 price_type=SALE 的 max 售价 vs CSV定价上限"],
        ["", ""],
        ["结果", "行数"],
        ["CSV行", len(csv_rows)],
        ["AD未删行", len(ads)],
        ["逐行-通过", counts.get("通过", 0)],
        ["逐行-警告", counts.get("警告", 0)],
        ["逐行-失败", counts.get("失败", 0)],
        ["逐行-未对码", counts.get("未对码", 0)],
        ["逐行-跳过", counts.get("跳过", 0)],
        ["ACTIVE未对码", len(unmatched_active)],
        ["STOPPED仍启用", len(stopped_still_on)],
        ["等级不一致", len(lv_mismatch)],
        ["名称漂移", len(name_drift)],
        ["售价超上限", len(over_cap)],
        ["一对多", len(multi_hit)],
        ["AD有yb_no不在本CSV", len(extra_ad)],
    ]
    for r in notes:
        ws0.append(r)
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 80

    write_sheet(
        wb,
        "逐行对照",
        [
            "csv行号",
            "catalog_kind",
            "csv_status",
            "业务键",
            "project_code",
            "hilist_code",
            "csv名称",
            "csv等级",
            "csv计价单位",
            "csv定价上限",
            "结果",
            "问题",
            "AD_id",
            "AD_bus_no",
            "AD_yb_no",
            "AD名称",
            "AD状态",
            "AD等级",
            "AD单位code",
            "AD单位显示",
            "售价min",
            "售价max",
            "overlay可用",
            "overlay止期",
            "命中AD数",
        ],
        detail,
        fail_col=10,
    )
    write_sheet(
        wb,
        "ACTIVE未对码",
        ["csv行号", "kind", "业务键", "project_code", "hilist_code", "名称", "等级", "单位", "上限", "地方编码"],
        unmatched_active,
    )
    write_sheet(
        wb,
        "STOPPED仍启用",
        ["csv行号", "kind", "业务键", "名称", "AD_id", "bus_no", "yb_no", "AD状态", "overlay可用", "overlay止期"],
        stopped_still_on,
    )
    write_sheet(
        wb,
        "等级不一致",
        ["csv行号", "kind", "业务键", "名称", "csv等级", "AD等级", "AD_id", "AD_yb_no"],
        lv_mismatch,
    )
    write_sheet(
        wb,
        "名称漂移",
        ["csv行号", "kind", "业务键", "csv名称", "AD名称", "AD_id", "bus_no"],
        name_drift,
    )
    write_sheet(
        wb,
        "售价超上限",
        ["csv行号", "kind", "业务键", "名称", "定价上限", "AD售价max", "AD_id"],
        over_cap,
    )
    write_sheet(
        wb,
        "一对多命中",
        ["csv行号", "kind", "业务键", "名称", "AD_ids", "bus_nos"],
        multi_hit,
    )
    write_sheet(
        wb,
        "AD有码不在目录",
        ["AD_id", "bus_no", "yb_no", "名称", "状态", "等级"],
        extra_ad,
    )
    wb.save(OUT_PATH)
    print("WROTE", OUT_PATH)
    print("csv", len(csv_rows), "ad", len(ads))
    print(dict(counts))
    print(
        "unmatched_active",
        len(unmatched_active),
        "stopped_on",
        len(stopped_still_on),
        "lv",
        len(lv_mismatch),
        "name",
        len(name_drift),
        "overcap",
        len(over_cap),
        "extra",
        len(extra_ad),
    )


if __name__ == "__main__":
    main()
