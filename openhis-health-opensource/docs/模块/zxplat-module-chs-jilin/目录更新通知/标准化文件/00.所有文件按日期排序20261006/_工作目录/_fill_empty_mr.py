# -*- coding: utf-8 -*-
"""用国基诊疗目录表补 whale-training 中为空的病案结算类型。不改已有值，不改财务类型。"""
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import psycopg2

XLSX = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\数据迁移相关\服务目录\国基医院新系统诊疗目录表.xlsx"
OUT = r"D:\0.product\03.ZXPlat\zxplat\docs\projects\project-H22010202493\数据迁移相关\服务目录\病案结算类型补齐结果.xlsx"
MR = {
    "一般医疗服务费": "1", "一般治疗操作费": "2", "护理费": "3", "综合医疗服务类其他费用": "4",
    "病理诊断费": "5", "实验室诊断费": "6", "影像学诊断费": "7", "临床诊断项目费": "8",
    "非手术治疗项目费": "9", "非手术治疗项目费 (目录节点)": "9",
    "临床物理治疗费": "10", "手术治疗费": "11", "手术治疗费 (目录节点)": "11",
    "麻醉费": "12", "手术费": "13", "康复费": "14", "中医治疗费": "15",
    "中医骨伤": "15", "中医特殊治疗": "15", "中医推拿治疗": "15", "针刺与灸法": "15", "中医外治": "15",
    "西药费": "16", "抗菌药物费用": "17", "中成药费": "18", "中药饮片费": "19", "血费": "20",
    "白蛋白类制品费": "21", "球蛋白类制品费": "22", "凝血因子类制品费": "23", "细胞因子类制品费": "24",
    "检查用一次性医用材料费": "25", "治疗用一次性医用材料费": "26", "手术用一次性医用材料费": "27",
    "其他费": "28", "其他费用": "28",
}
CODE_LABEL = {code: label for label, code in MR.items() if "目录节点" not in label and label not in (
    "中医骨伤", "中医特殊治疗", "中医推拿治疗", "针刺与灸法", "中医外治", "其他费用")}


def text(value):
    return str(value or "").strip()


def code_key(value):
    return text(value).replace(" ", "")


def load_indexes():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    by_name = defaultdict(set)
    by_code = defaultdict(set)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        header = [text(cell) for cell in rows[0]]
        name_cols = [i for i, head in enumerate(header) if head in ("*诊疗名称", "项目", "项目名称")]
        code_cols = [i for i, head in enumerate(header) if head in ("医保编号", "医保编码", "项目代码", "编码")]
        mr_cols = [i for i, head in enumerate(header) if head in ("病案结算类型", "病案结算项")]
        if not mr_cols:
            continue
        for row in rows[1:]:
            codes = set()
            for index in mr_cols:
                if index < len(row):
                    label = text(row[index])
                    if label in MR:
                        codes.add(MR[label])
            if not codes:
                continue
            for index in name_cols:
                if index < len(row):
                    name = text(row[index])
                    if name and not name.startswith("[示例]"):
                        by_name[name].update(codes)
            for index in code_cols:
                if index >= len(row):
                    continue
                raw = code_key(row[index])
                if not raw or raw.startswith("[示例]"):
                    continue
                by_code[raw].update(codes)
                head = raw.split("-")[0]
                if head != raw:
                    by_code[head].update(codes)
    wb.close()
    return by_name, by_code


def unique(codes):
    if codes and len(codes) == 1:
        return next(iter(codes))
    return None


def main():
    by_name, by_code = load_indexes()
    conn = psycopg2.connect(host="192.168.1.1", port=5433, user="postgres", password="root", dbname="whale-training")
    cur = conn.cursor()
    cur.execute("SELECT current_database()")
    if cur.fetchone()[0] != "whale-training":
        raise SystemExit("拒绝：不是 whale-training")
    cur.execute("""
        SELECT id, bus_no, name, coalesce(yb_no, '')
        FROM wkf_activity_definition
        WHERE is_deleted = false
          AND (mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '')
        ORDER BY bus_no
    """)
    empty = cur.fetchall()
    updates = []
    left = []
    for ad_id, bus, name, yb in empty:
        yb_key = code_key(yb)
        chosen = unique(by_name.get(name))
        how = "名称"
        if chosen is None and yb_key:
            chosen = unique(by_code.get(yb_key))
            how = "医保编码"
        if chosen is None and yb_key:
            chosen = unique(by_code.get(yb_key.split("-")[0]))
            how = "医保编码前段"
        if chosen:
            updates.append((ad_id, bus, name, yb_key, chosen, how))
        else:
            left.append((bus, name, yb_key))
    print("empty", len(empty), "update", len(updates), "left", len(left))
    cur.executemany(
        """
        UPDATE wkf_activity_definition
        SET mr_settlement_type_code = %s,
            updated_at = CURRENT_TIMESTAMP,
            updated_by = 1
        WHERE id = %s
          AND is_deleted = false
          AND (mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '')
        """,
        [(code, ad_id) for ad_id, _bus, _name, _yb, code, _how in updates],
    )
    print("updated_rows", cur.rowcount)
    conn.commit()
    cur.execute("""
        SELECT
          SUM(CASE WHEN mr_settlement_type_code IS NULL OR btrim(mr_settlement_type_code) = '' THEN 1 ELSE 0 END)
        FROM wkf_activity_definition WHERE is_deleted = false
    """)
    print("mr still empty", cur.fetchone()[0])
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "已补齐"
    ws.append(["业务编码", "项目名称", "医保编码", "病案结算类型", "编码", "匹配方式"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
    for ad_id, bus, name, yb, code, how in updates:
        ws.append([bus, name, yb, CODE_LABEL.get(code, code), code, how])
    left_sheet = wb.create_sheet("文件中没有")
    left_sheet.append(["业务编码", "项目名称", "医保编码"])
    for row in left:
        left_sheet.append(list(row))
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
