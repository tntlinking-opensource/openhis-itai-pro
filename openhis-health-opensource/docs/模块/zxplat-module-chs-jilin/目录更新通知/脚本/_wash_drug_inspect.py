# -*- coding: utf-8 -*-
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

base = Path(r"D:\0.product\02.Whale\whale-doc\whale-module-chs-jilin\目录更新通知")
targets = [
    "附件：药品数据更新目录（2026.1.26、2.4）.xlsx",
    "附件：药品数据更新目录（2025.9.5、9.19）.xlsx",
    "_附件：药品数据更新目录（2025.10.30、11.10、11.14）.xlsx",
    "附件：药品数据更新目录（2025.3.7) .xlsx",
    "附件：药品数据更新目录（2025.1.26).xlsx",
    "附件：药品数据更新目录（Y2025.5.13）.xlsx",
    "附件：药品数据更新目录（X001_20251224000000_A）.xlsx",
    "附件：药品数据更新目录（2025.01.01）.xlsx",
    "药品数据更新目录（2024.8.23）.xlsx",
]

for t in targets:
    p = base / t
    print(f"\n==== {t} | exists={p.exists()} size={p.stat().st_size if p.exists() else 0} ====")
    if not p.exists():
        continue
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
        print("  sheets:", wb.sheetnames)
        for sn in wb.sheetnames[:6]:
            ws = wb[sn]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                rows.append([(c if c is None else str(c)[:40]) for c in row[:12]])
                if i >= 2:
                    break
            print(f"  -- {sn}")
            for r in rows:
                print("    ", r)
        if len(wb.sheetnames) > 6:
            print(f"  ... +{len(wb.sheetnames)-6} more sheets")
        wb.close()
    except Exception as e:
        print("  ERROR:", type(e).__name__, e)
