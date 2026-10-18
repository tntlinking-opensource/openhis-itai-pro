# -*- coding: utf-8 -*-
"""提示词 1A：局发附件 → 标准落地包 + 归档。"""
from __future__ import annotations

import csv
import json
import re
import shutil
import traceback
from collections import Counter
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"D:\0.product\02.Whale\whale-doc\whale-module-chs-jilin\目录更新通知")
STD_ROOT = BASE / "标准化文件"
RAW_ROOT = BASE / "原始文件"
FAIL_ROOT = RAW_ROOT / "_无法转换"

TARGETS = [
    "2024.10.25中药饮片（林下山参）.xlsx",
    "附件：中药饮片数据更新目录（2025.5.21）.xlsx",
    "附件：耗材数据更新目录（2026.5.9）.xlsx",
    "附件：耗材数据更新目录（2026.5.9） (1).xlsx",
    "附件：耗材数据更新（2025.7.23、2025.8.7）.xlsx",
    "附件：耗材数据更新目录（2025.9.13）.xlsx",
    "附件：耗材数据更新目录（2025.4.24）.xlsx",
    "附件：1.耗材目录更新目录（2025.10.13）.xlsx",
    "附件：耗材数据更新目录（2025.6.28）.xlsx",
    "耗材数据更新（H2024.7.1）.xlsx",
    "附件：院内制剂目录更新（2025.12.30）.xlsx",
    "附件：中药配方颗粒(2025.3.7) .xlsx",
    "附件：中药配方颗粒（2025.08.11）.xlsx",
    "中药配方颗粒（2024.11.15）.xlsx",
    "中药配方颗粒更新目录（2024.8.23）.xlsx",
]

HILIST_COLS = [
    "医疗目录编码",
    "药品代码",
    "目录编码",
    "国家目录编码",
    "中药配方颗粒代码",
    "制剂代码",
]
NAME_COLS = [
    "医保目录名称",
    "中草药名称",
    "耗材名称",
    "中药配方颗粒名称",
    "制剂名称",
    "药品名称",
    "项目名称",
]
CHANGE_COLS = ["本期变更", "变更标记", "本次变更"]
PRICE_KEYS = ["医保目录定价上限金额", "医保目录超限处理方式"]

CSV_FIELDS = [
    "action",
    "hilist_code",
    "project_code",
    "hilist_name",
    "begin_date",
    "end_date",
    "change_fields",
    "chrgitm_lv",
    "list_type",
    "source_sheet",
    "source_row",
    "ext_json",
]

HEADER_HINTS = set(HILIST_COLS + NAME_COLS + ["开始日期", "终止日期", "收费项目等级", "本期变更"])


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def to_date(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s or s in ("--", "None", "null"):
        return ""
    # 2022-02-28 00:00:00.0
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    # excel serial unlikely via openpyxl data_only
    return ""


def map_chrgitm_lv(v) -> str:
    if v is None or v == "":
        return ""
    s = cell_str(v)
    if not s:
        return ""
    low = s.lower()
    if s in ("1", "01", "甲"):
        return "1"
    if s in ("2", "02", "乙"):
        return "2"
    if s in ("3", "03", "丙", "自费") or "自费" in s or s.startswith("丙"):
        return "3"
    # 纯数字 1/2/3
    if re.fullmatch(r"0*[123]", s):
        return str(int(s))
    return ""


def infer_category(filename: str) -> str:
    if "配方颗粒" in filename:
        return "FORMULA"
    if "饮片" in filename:
        return "HERBAL"
    if "制剂" in filename:
        return "PREP"
    if "耗材" in filename:
        return "DEVICE"
    return "OTHER"


def default_list_type(category: str) -> str:
    return {
        "DRUG": "101",
        "SERVICE": "201",
        "SERVICE_F": "201",
        "DEVICE": "301",
        "HERBAL": "102",
        "FORMULA": "106",
        "PREP": "103",
    }.get(category, "")


def parse_notice_date(filename: str) -> str | None:
    # H2024.7.1 / 2025.08.11 / 2024.10.25 / 2025.7.23、2025.8.7
    dates = re.findall(r"(?:H)?(\d{4})\.(\d{1,2})\.(\d{1,2})", filename)
    if not dates:
        return None
    y, m, d = dates[-1]  # 多日期取末个作为通知日代表
    return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"


def sheet_action(sheet_name: str) -> str | None:
    n = sheet_name.strip()
    if n.startswith("省变更") or "省变更" in n[:6]:
        return "PROVINCE_MODIFY"
    if n.startswith("新增"):
        return "ADD"
    if n.startswith("停用"):
        return "STOP"
    if n.startswith("变更"):
        return "MODIFY"
    # 可识别：院内制剂变更 / xxx变更
    if "变更" in n and not n.startswith("新增") and not n.startswith("停用"):
        return "MODIFY"
    return None


def find_header_row(rows: list) -> tuple[int, list[str]] | None:
    for i, row in enumerate(rows):
        vals = [cell_str(c) for c in row]
        non_empty = [v for v in vals if v]
        if not non_empty:
            continue
        hit = sum(1 for v in non_empty if v in HEADER_HINTS)
        # 也认常见编码列名
        if hit >= 2 or any(v in HILIST_COLS for v in non_empty):
            # normalize: strip empties at end for indexing but keep positions
            return i, vals
    return None


def pick_from(row_map: dict, cols: list[str]) -> str:
    for c in cols:
        if c in row_map and row_map[c]:
            return row_map[c]
    return ""


def json_safe(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return to_date(v) or cell_str(v)
    if isinstance(v, (int, float, bool)):
        return v
    s = cell_str(v)
    return s if s != "" else None


def unique_dest(dest_dir: Path, filename: str) -> Path:
    dest = dest_dir / filename
    if not dest.exists():
        return dest
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    ts = datetime.now().strftime("%Y%m%d%H%M%S")
    return dest_dir / f"{stem}_{ts}{suffix}"


def safe_stem(src: Path) -> str:
    """Windows 目录名不能以空格/点结尾；去掉文件名 stem 首尾空白。"""
    return src.stem.strip().rstrip(". ")


def wash_one(src: Path) -> dict:
    stem = safe_stem(src)
    out_dir = STD_ROOT / stem
    category = infer_category(src.name)
    notice_date = parse_notice_date(src.name)
    wash_notes = []
    if stem != src.stem:
        wash_notes.append(f"标准目录名已trim：{src.stem!r}→{stem!r}")

    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except Exception as e:
        return fail_move(src, f"无法打开 xlsx：{type(e).__name__}: {e}")

    lines: list[dict] = []
    tried_sheets = []
    try:
        for sn in wb.sheetnames:
            action = sheet_action(sn)
            tried_sheets.append(f"{sn}->{action or '?'}")
            if not action:
                wash_notes.append(f"跳过无法识别动作的 Sheet：{sn}")
                continue
            ws = wb[sn]
            rows = list(ws.iter_rows(values_only=True))
            hdr = find_header_row(rows)
            if not hdr:
                wash_notes.append(f"Sheet {sn} 未找到表头")
                continue
            hdr_idx, headers = hdr
            # pad headers for blank names
            headers = [h if h else f"__empty_{i}" for i, h in enumerate(headers)]
            for r_i, row in enumerate(rows[hdr_idx + 1 :], start=hdr_idx + 2):
                vals = list(row)
                if all(v is None or str(v).strip() == "" for v in vals):
                    continue
                # extend
                while len(vals) < len(headers):
                    vals.append(None)
                row_map = {}
                for h, v in zip(headers, vals):
                    if h.startswith("__empty_"):
                        continue
                    row_map[h] = cell_str(v) if not isinstance(v, (datetime, date)) else to_date(v) or cell_str(v)
                    # keep raw for dates specifically
                # override dates properly from original
                for di, h in enumerate(headers):
                    if h in ("开始日期", "终止日期"):
                        row_map[h] = to_date(vals[di])

                hilist_code = pick_from(row_map, HILIST_COLS)
                if not hilist_code:
                    continue
                hilist_name = pick_from(row_map, NAME_COLS)
                change_fields = pick_from(row_map, CHANGE_COLS)
                chrgitm_lv = map_chrgitm_lv(row_map.get("收费项目等级", ""))
                list_type = cell_str(row_map.get("目录类别", "")) or default_list_type(category)
                begin_date = row_map.get("开始日期", "")
                end_date = row_map.get("终止日期", "")

                mapped = set(
                    HILIST_COLS
                    + NAME_COLS
                    + CHANGE_COLS
                    + ["开始日期", "终止日期", "收费项目等级", "目录类别"]
                )
                ext = {}
                for h, v in row_map.items():
                    if h in mapped:
                        # price keys still go to ext
                        if h in PRICE_KEYS and v != "":
                            ext[h] = json_safe(v)
                        continue
                    if v != "":
                        ext[h] = json_safe(v)
                for pk in PRICE_KEYS:
                    if pk in row_map and row_map[pk] != "" and pk not in ext:
                        ext[pk] = json_safe(row_map[pk])

                lines.append(
                    {
                        "action": action,
                        "hilist_code": hilist_code,
                        "project_code": "",
                        "hilist_name": hilist_name,
                        "begin_date": begin_date,
                        "end_date": end_date,
                        "change_fields": change_fields,
                        "chrgitm_lv": chrgitm_lv,
                        "list_type": list_type,
                        "source_sheet": sn,
                        "source_row": str(r_i),
                        "ext_json": json.dumps(ext, ensure_ascii=False, separators=(",", ":")),
                    }
                )
    finally:
        wb.close()

    if not lines:
        return fail_move(
            src,
            "无可转换明细行。"
            f" Sheets={tried_sheets}; notes={wash_notes}",
        )

    out_dir.mkdir(parents=True, exist_ok=True)
    # clean previous outputs if re-run
    for fn in ("notice.meta.json", "notice.lines.csv"):
        p = out_dir / fn
        if p.exists():
            p.unlink()

    action_counts = Counter(x["action"] for x in lines)
    note_parts = [
        f"CATALOG_DELTA/{category}",
        f"行数={len(lines)}",
        "actions=" + ",".join(f"{k}:{v}" for k, v in sorted(action_counts.items())),
    ]
    if wash_notes:
        note_parts.append("；".join(wash_notes))
    if category == "DEVICE":
        note_parts.append("hilist_code←国家目录编码；限价列进ext_json")
    elif category == "HERBAL":
        note_parts.append("hilist_code←医疗目录编码")
    elif category == "FORMULA":
        note_parts.append("hilist_code←中药配方颗粒代码")
    elif category == "PREP":
        note_parts.append("hilist_code←制剂代码；Sheet名含变更→MODIFY")

    meta = {
        "doc_kind": "CATALOG_DELTA",
        "category": category,
        "title": stem,
        "source_filename": src.name,
        "notice_date": notice_date,
        "effective_date": None,
        "wash_note": "；".join(note_parts),
        "wash_tool": "cursor-ai",
    }
    (out_dir / "notice.meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    with open(out_dir / "notice.lines.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        for row in lines:
            w.writerow(row)

    RAW_ROOT.mkdir(parents=True, exist_ok=True)
    dest = unique_dest(RAW_ROOT, src.name)
    shutil.move(str(src), str(dest))

    return {
        "result": "OK",
        "source": src.name,
        "std_dir": str(out_dir.relative_to(BASE)),
        "moved_to": str(dest.relative_to(BASE)),
        "rows": len(lines),
        "actions": dict(action_counts),
        "category": category,
        "doc_kind": "CATALOG_DELTA",
        "notice_date": notice_date,
        "remark": meta["wash_note"],
    }


def fail_move(src: Path, reason: str) -> dict:
    """失败：原件进 原始文件/_无法转换/{stem}/，并写 reason.txt。"""
    FAIL_ROOT.mkdir(parents=True, exist_ok=True)
    stem = safe_stem(src)
    sub = FAIL_ROOT / stem
    # 同名目录已存在则加时间后缀
    if sub.exists() and any(sub.iterdir()):
        sub = FAIL_ROOT / f"{stem}_{datetime.now().strftime('%Y%m%d%H%M%S')}"
    sub.mkdir(parents=True, exist_ok=True)
    dest = sub / src.name
    shutil.move(str(src), str(dest))
    rp = sub / "reason.txt"
    rp.write_text(reason + "\n", encoding="utf-8")
    return {
        "result": "FAIL",
        "source": src.name,
        "std_dir": "",
        "moved_to": str(dest.relative_to(BASE)),
        "rows": 0,
        "actions": {},
        "category": infer_category(src.name),
        "doc_kind": "",
        "notice_date": parse_notice_date(src.name),
        "remark": reason,
        "reason_file": str(rp.relative_to(BASE)),
    }


def main():
    STD_ROOT.mkdir(parents=True, exist_ok=True)
    RAW_ROOT.mkdir(parents=True, exist_ok=True)
    results = []
    for name in TARGETS:
        src = BASE / name
        print("=" * 60)
        print("PROCESSING:", name)
        if not src.exists():
            # maybe already moved
            alt = RAW_ROOT / name
            fail_alt = FAIL_ROOT / name
            if alt.exists():
                rec = {
                    "result": "SKIP",
                    "source": name,
                    "std_dir": str((STD_ROOT / Path(name).stem).relative_to(BASE))
                    if (STD_ROOT / Path(name).stem).exists()
                    else "",
                    "moved_to": str(alt.relative_to(BASE)),
                    "rows": "",
                    "actions": {},
                    "remark": "根目录已不存在，位于原始文件/",
                }
            elif fail_alt.exists():
                rec = {
                    "result": "SKIP",
                    "source": name,
                    "std_dir": "",
                    "moved_to": str(fail_alt.relative_to(BASE)),
                    "rows": 0,
                    "actions": {},
                    "remark": "根目录已不存在，位于_无法转换/",
                }
            else:
                rec = {
                    "result": "FAIL",
                    "source": name,
                    "std_dir": "",
                    "moved_to": "",
                    "rows": 0,
                    "actions": {},
                    "remark": "根目录未找到该文件",
                }
            results.append(rec)
            print(json.dumps(rec, ensure_ascii=False, indent=2))
            continue
        try:
            rec = wash_one(src)
        except Exception as e:
            # unexpected — try fail archive
            try:
                rec = fail_move(src, f"洗数异常：{type(e).__name__}: {e}\n{traceback.format_exc()}")
            except Exception as e2:
                rec = {
                    "result": "FAIL",
                    "source": name,
                    "std_dir": "",
                    "moved_to": "",
                    "rows": 0,
                    "actions": {},
                    "remark": f"洗数异常且归档失败：{e} / {e2}",
                }
        results.append(rec)
        print(json.dumps(rec, ensure_ascii=False, indent=2))

    summary_path = BASE / "_wash_1a_batch_summary.json"
    summary_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print("SUMMARY ->", summary_path)


if __name__ == "__main__":
    main()
