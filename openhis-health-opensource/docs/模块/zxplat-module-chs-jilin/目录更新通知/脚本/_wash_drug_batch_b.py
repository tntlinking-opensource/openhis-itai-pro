# -*- coding: utf-8 -*-
"""Wash drug catalog delta xlsx -> §6.0 standard packs (batch B)."""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
from collections import Counter
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"D:\0.product\02.Whale\whale-doc\whale-module-chs-jilin\目录更新通知")
STD_ROOT = BASE / "标准化文件"
RAW_ROOT = BASE / "原始文件"
FAIL_ROOT = RAW_ROOT / "_无法转换"

TARGETS = [
    "附件：药品数据更新目录（2026.1.26、2.4）.xlsx",
    "附件：药品数据更新目录（2025.9.5、9.19）.xlsx",
    "_附件：药品数据更新目录（2025.10.30、11.10、11.14）.xlsx",
    "附件：药品数据更新目录（2025.8.20、8.28、9.02）.xlsx",
    "附件：药品数据更新目录（2026.3.19、3.27、4.2、4.10）.xlsx",
    "附件：药品数据更新目录（2026.2.10、2.12、2.28）.xlsx",
    "附件：药品数据更新目录（2026.1.4、1.12、1.15）.xlsx",
    "附件：药品数据更新目录（2025.3.7) .xlsx",
    "附件：药品数据更新目录（2025.9.26、9.30、10.14、10.21）.xlsx",
    "附件：药品数据更新目录（2025.1.26).xlsx",
    "附件：药品数据更新目录（2025.7.24、7.30、8.4、8.8）.xlsx",
    "附件：药品数据更新目录（2025.1.13）.xlsx",
    "附件：药品数据更新目录（2025.11.24、12.01、12.10、12.17、12.19）.xlsx",
    "附件：药品数据更新目录（Y2025.5.13）.xlsx",
    "附件：药品数据更新目录（X001_20251224000000_A）.xlsx",
    "附件：药品数据更新目录（2024.12.30）.xlsx",
    "附件：药品数据更新目录（2024.12.17）.xlsx",
    "附件：药品数据更新目录（2025.1.2）.xlsx",
    "附件：药品数据更新目录（2025.01.01）.xlsx",
    "药品数据更新目录（2024.8.23）.xlsx",
]

CSV_HEADERS = [
    "action",
    "hilist_code",
    "project_code",
    "hilist_name",
    "begin_date",
    "end_date",
    "change_fields",
    "chrgitm_lv",
    "list_type",
    "source_sheet",
    "source_row",
    "ext_json",
]

CODE_KEYS = ("医疗目录编码", "药品代码", "目录编码", "国家目录编码", "中药配方颗粒代码")
NAME_KEYS = ("注册名称", "药品名称", "药品商品名", "项目名称", "通用名", "药品通用名")
BEGIN_KEYS = ("开始时间", "开始日期")
END_KEYS = ("终止日期", "终止时间", "结束日期")
CHANGE_KEYS = ("本期变更",)
LV_KEYS = ("收费项目等级", "国家医保药品目录甲乙类标识")
LIST_TYPE_KEYS = ("目录类别",)
CORE_KEYS = set(CODE_KEYS + NAME_KEYS + BEGIN_KEYS + END_KEYS + CHANGE_KEYS + LV_KEYS + LIST_TYPE_KEYS)

SHEET_ACTIONS = [
    ("省变更", "PROVINCE_MODIFY"),
    ("新增", "ADD"),
    ("停用", "STOP"),
    ("变更", "MODIFY"),
]


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.lower() in ("none", "null", "nan", "--", "—", "-"):
        return ""
    return s


def normalize_header(h) -> str:
    return cell_str(h).replace("\n", "").replace("\r", "").strip()


def parse_date(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = cell_str(v)
    if not s:
        return ""
    s = s.replace("/", "-").replace(".", "-")
    # datetime-like
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return ""
    # compact yyyymmdd
    m = re.match(r"^(\d{4})(\d{2})(\d{2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return ""
    return ""


def map_chrgitm_lv(v) -> str:
    s = cell_str(v)
    if not s:
        return ""
    if s in ("1", "01", "甲"):
        return "1"
    if s in ("2", "02", "乙"):
        return "2"
    if s in ("3", "03", "丙", "自费"):
        return "3"
    # numeric with leading zeros already handled; try int
    try:
        n = int(float(s))
        if n in (1, 2, 3):
            return str(n)
    except ValueError:
        pass
    return ""


def sheet_action(sheet_name: str) -> str | None:
    for prefix, action in SHEET_ACTIONS:
        if sheet_name.startswith(prefix):
            return action
    return None


def parse_notice_dates(filename: str) -> list[str]:
    stem = Path(filename).stem
    dates: list[str] = []
    # X001_YYYYMMDDHHMMSS
    m = re.search(r"X001_(\d{8})", stem)
    if m:
        d = parse_date(m.group(1))
        if d:
            dates.append(d)
    # Y2025.5.13 or Y2025.05.13
    m = re.search(r"Y(\d{4})\.(\d{1,2})\.(\d{1,2})", stem)
    if m:
        d = parse_date(f"{m.group(1)}-{m.group(2)}-{m.group(3)}")
        if d:
            dates.append(d)
    # year-only scope in paren: （2026.1.26、2.4） / （2025.01.01）
    m = re.search(r"[（(]([^）)]+)[）)]", stem)
    if m:
        blob = m.group(1)
        # full dates with year
        for y, mo, d in re.findall(r"(\d{4})\.(\d{1,2})\.(\d{1,2})", blob):
            dd = parse_date(f"{y}-{mo}-{d}")
            if dd:
                dates.append(dd)
        # year prefix once, then md fragments like 2.4 / 12.01
        ym = re.search(r"(\d{4})\.(\d{1,2})\.(\d{1,2})", blob)
        if ym:
            year = ym.group(1)
            # fragments after first full date: 、2.4 or 、12.01
            rest = blob[ym.end() :]
            for mo, d in re.findall(r"[、,，]\s*(\d{1,2})\.(\d{1,2})", rest):
                dd = parse_date(f"{year}-{mo}-{d}")
                if dd:
                    dates.append(dd)
    # plain （2024.8.23） already covered
    # dedupe preserve order
    seen = set()
    out = []
    for d in dates:
        if d not in seen:
            seen.add(d)
            out.append(d)
    return out


def pick_value(row: dict, keys: tuple[str, ...]) -> str:
    for k in keys:
        if k in row and row[k]:
            return row[k]
    return ""


def row_to_ext(row: dict, used: set[str]) -> str:
    ext = {}
    for k, v in row.items():
        if not k or k in used or k in CORE_KEYS:
            continue
        if v == "":
            continue
        ext[k] = v
    # price-related already in CORE? 医保目录定价上限金额 etc go to ext via not-in-CORE
    return json.dumps(ext, ensure_ascii=False, separators=(",", ":"))


def header_index(headers: list[str]) -> dict[str, int]:
    idx = {}
    for i, h in enumerate(headers):
        if h and h not in idx:
            idx[h] = i
    return idx


def wash_file(src: Path) -> dict:
    stem_raw = src.stem  # keep prefix incl. _附件
    # Windows 目录名不能保留尾随空格/点
    stem = stem_raw.rstrip(" .")
    out_dir = STD_ROOT / stem
    dates = parse_notice_dates(src.name)
    notice_date = dates[-1] if dates else None

    wb = load_workbook(src, read_only=True, data_only=True)
    lines: list[dict] = []
    known_sheets = 0
    notes = []
    if stem != stem_raw:
        notes.append(f"标准目录名已rstrip尾随空格/点")

    try:
        for sn in wb.sheetnames:
            action = sheet_action(sn)
            if not action:
                notes.append(f"跳过未识别Sheet:{sn}")
                continue
            known_sheets += 1
            ws = wb[sn]
            headers = None
            header_map = None
            for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = list(row)
                if headers is None:
                    headers = [normalize_header(c) for c in vals]
                    # skip empty header rows
                    if not any(headers):
                        headers = None
                        continue
                    header_map = header_index(headers)
                    # require a code column
                    if not any(k in header_map for k in CODE_KEYS):
                        notes.append(f"Sheet无编码列:{sn}")
                        break
                    continue
                # build dict
                d = {}
                for h, i in header_map.items():
                    if i < len(vals):
                        d[h] = cell_str(vals[i])
                    else:
                        d[h] = ""
                code = pick_value(d, CODE_KEYS)
                if not code:
                    # skip blank / total rows
                    continue
                name = pick_value(d, NAME_KEYS)
                begin = parse_date(pick_value(d, BEGIN_KEYS) or None)
                # also try raw cell if pick returned non-empty but parse failed earlier
                if not begin:
                    for k in BEGIN_KEYS:
                        if k in d and d[k]:
                            begin = parse_date(d[k])
                            break
                end = ""
                for k in END_KEYS:
                    if k in d and d[k]:
                        end = parse_date(d[k])
                        break
                # preserve unparseable end in ext by keeping original if parse failed
                raw_end = pick_value(d, END_KEYS)
                change = pick_value(d, CHANGE_KEYS)
                lv = ""
                for k in LV_KEYS:
                    if k in d and d[k]:
                        lv = map_chrgitm_lv(d[k])
                        if lv:
                            break
                list_type = pick_value(d, LIST_TYPE_KEYS) or "101"
                # normalize list_type digits
                if list_type:
                    try:
                        list_type = str(int(float(list_type)))
                    except ValueError:
                        pass

                used = set()
                for k in CODE_KEYS + NAME_KEYS + BEGIN_KEYS + END_KEYS + CHANGE_KEYS + LV_KEYS + LIST_TYPE_KEYS:
                    used.add(k)
                # if end date unparseable, stash original
                ext_extra = {}
                if raw_end and not end:
                    ext_extra["_raw_终止日期"] = raw_end

                ext_obj = convert_ext(d, used, ext_extra)

                lines.append(
                    {
                        "action": action,
                        "hilist_code": code,
                        "project_code": "",
                        "hilist_name": name,
                        "begin_date": begin,
                        "end_date": end,
                        "change_fields": change,
                        "chrgitm_lv": lv,
                        "list_type": list_type,
                        "source_sheet": sn,
                        "source_row": str(r_i),
                        "ext_json": json.dumps(ext_obj, ensure_ascii=False, separators=(",", ":")),
                    }
                )
    finally:
        wb.close()

    if known_sheets == 0:
        raise ValueError("无识别的目录更新Sheet（需新增/停用/变更/省变更前缀）")
    if not lines:
        raise ValueError("识别到Sheet但无有效数据行")

    out_dir.mkdir(parents=True, exist_ok=True)
    meta = {
        "doc_kind": "CATALOG_DELTA",
        "category": "DRUG",
        "title": stem_raw.rstrip(),
        "source_filename": src.name,
        "notice_date": notice_date,
        "effective_date": None,
        "wash_note": "; ".join(
            [
                "药品目录更新xlsx→标准包",
                f"sheet动作按前缀映射ADD/STOP/MODIFY/PROVINCE_MODIFY",
                f"hilist_code←医疗目录编码|药品代码",
                f"list_type默认101或取目录类别",
                f"文件名日期:{','.join(dates) if dates else '无'}",
            ]
            + notes
        ),
        "wash_tool": "cursor-ai",
    }
    (out_dir / "notice.meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    csv_path = out_dir / "notice.lines.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_HEADERS, extrasaction="ignore")
        w.writeheader()
        for line in lines:
            w.writerow(line)

    counts = Counter(l["action"] for l in lines)
    return {
        "ok": True,
        "out_dir": str(out_dir),
        "rows": len(lines),
        "counts": dict(counts),
        "notice_date": notice_date,
        "meta": meta,
    }


def convert_ext(row: dict, used: set[str], extra: dict) -> dict:
    ext = dict(extra)
    for k, v in row.items():
        if not k or k in used or k in CORE_KEYS:
            continue
        if v == "":
            continue
        ext[k] = v
    return ext


def unique_dest(dir_path: Path, name: str) -> Path:
    dest = dir_path / name
    if not dest.exists():
        return dest
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    p = Path(name)
    return dir_path / f"{p.stem}_{stamp}{p.suffix}"


def move_ok(src: Path) -> Path:
    RAW_ROOT.mkdir(parents=True, exist_ok=True)
    dest = unique_dest(RAW_ROOT, src.name)
    shutil.move(str(src), str(dest))
    return dest


def move_fail(src: Path, reason: str) -> Path:
    FAIL_ROOT.mkdir(parents=True, exist_ok=True)
    dest = unique_dest(FAIL_ROOT, src.name)
    shutil.move(str(src), str(dest))
    (FAIL_ROOT / f"{dest.stem}.reason.txt").write_text(
        f"file: {src.name}\n{reason}\n", encoding="utf-8"
    )
    return dest


def main():
    STD_ROOT.mkdir(parents=True, exist_ok=True)
    RAW_ROOT.mkdir(parents=True, exist_ok=True)
    report = []
    for name in TARGETS:
        src = BASE / name
        print(f"\n######## {name} ########", flush=True)
        if not src.exists():
            print("SKIP: not in root", flush=True)
            report.append((name, "SKIP", "", "", "根目录不存在(可能已被处理)"))
            continue
        try:
            result = wash_file(src)
            dest = move_ok(src)
            cnt = result["counts"]
            cnt_str = ",".join(f"{k}:{v}" for k, v in sorted(cnt.items()))
            print(
                f"OK rows={result['rows']} {cnt_str} notice_date={result['notice_date']} -> {result['out_dir']}",
                flush=True,
            )
            print(f"MOVED -> {dest}", flush=True)
            report.append((name, "OK", Path(result["out_dir"]).name, cnt_str, f"rows={result['rows']}"))
        except Exception as e:
            reason = f"{type(e).__name__}: {e}"
            print(f"FAIL {reason}", flush=True)
            try:
                dest = move_fail(src, reason)
                print(f"MOVED FAIL -> {dest}", flush=True)
            except Exception as e2:
                print(f"MOVE FAIL ERROR: {e2}", flush=True)
            report.append((name, "FAIL", "", "", reason))

    print("\n======= SUMMARY =======")
    for name, status, std, cnt, note in report:
        print(f"{status}|{std}|{cnt}|{note}|{name}")

    summary_path = BASE / "_wash_drug_batch_b_summary.json"
    summary_path.write_text(
        json.dumps(
            [{"file": a, "status": b, "std": c, "actions": d, "note": e} for a, b, c, d, e in report],
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print("summary ->", summary_path)


if __name__ == "__main__":
    main()
