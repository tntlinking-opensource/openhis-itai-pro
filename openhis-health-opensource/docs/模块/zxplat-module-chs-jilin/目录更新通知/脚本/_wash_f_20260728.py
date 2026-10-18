# -*- coding: utf-8 -*-
"""提示词 1A：F服务目录5类工程师版（2026.7.28）.xlsx → 标准落地包。
原件已在 原始文件/，仅产出标准包，不再 move。
"""
from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"D:\0.product\02.Whale\whale-doc\whale-module-chs-jilin\目录更新通知")
RAW = BASE / "原始文件"
STD_ROOT = BASE / "标准化文件"
SRC_NAME = "F服务目录5类工程师版（2026.7.28）.xlsx"

CSV_FIELDS = [
    "action",
    "hilist_code",
    "project_code",
    "hilist_name",
    "begin_date",
    "end_date",
    "change_fields",
    "chrgitm_lv",
    "list_type",
    "source_sheet",
    "source_row",
    "ext_json",
]

CORE_SKIP = {
    "项目代码",
    "项目名称",
    "医疗目录编码",
    "国家目录编码",
    "医保目录名称",
    "医疗服务项目名称",
    "开始日期",
    "终止日期",
    "收费项目等级",
    "目录类别",
    "本期变更",
    "变更标记",
    "本次变更",
}


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].lstrip("-").isdigit():
        s = s[:-2]
    return s


def to_date(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s or s in ("--", "None", "null"):
        return ""
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return ""


def map_chrgitm_lv(v) -> str:
    s = cell_str(v)
    if not s:
        return ""
    if s in ("1", "01", "甲"):
        return "1"
    if s in ("2", "02", "乙"):
        return "2"
    if s in ("3", "03", "丙", "自费") or "自费" in s or s.startswith("丙"):
        return "3"
    if re.fullmatch(r"0*[123]", s):
        return str(int(s))
    return ""


def sheet_action(sheet_name: str) -> str | None:
    n = sheet_name.strip()
    if n.startswith("省变更") or "省变更" in n[:6]:
        return "PROVINCE_MODIFY"
    if n.startswith("新增"):
        return "ADD"
    if n.startswith("停用"):
        return "STOP"
    if n.startswith("变更"):
        return "MODIFY"
    if "变更" in n:
        return "MODIFY"
    return None


def json_safe(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return to_date(v) or cell_str(v)
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v == int(v):
            return int(v)
        return v
    s = cell_str(v)
    return s if s != "" else None


def normalize_change_fields(v: str) -> str:
    s = cell_str(v)
    if not s or s == "--":
        return ""
    s = s.replace("、", ";").replace("，", ";").replace(",", ";")
    parts = [p.strip() for p in s.split(";") if p.strip()]
    return ";".join(parts)


def build_row_maps(headers: list[str], vals: list):
    """支持重复列名：保留首次非空；同时提供 raw 列表供日期取原始类型。"""
    row_map: dict[str, str] = {}
    raw_map: dict[str, object] = {}
    for h, v in zip(headers, vals):
        if not h:
            continue
        raw = v
        if h in ("开始日期", "终止日期"):
            s = to_date(v)
        else:
            s = cell_str(v)
        if h not in row_map or (not row_map[h] and s):
            row_map[h] = s
            raw_map[h] = raw
        elif s and not row_map[h]:
            row_map[h] = s
            raw_map[h] = raw
    return row_map, raw_map


def useful_ext_value(v) -> bool:
    if v is None:
        return False
    s = cell_str(v) if not isinstance(v, (int, float, bool)) else str(v)
    return s not in ("", "--", "None", "null")


def main():
    src = RAW / SRC_NAME
    if not src.exists():
        raise SystemExit(f"missing source: {src}")

    stem = src.stem.strip().rstrip(". ")
    out_dir = STD_ROOT / stem
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(src, read_only=True, data_only=True)
    lines: list[dict] = []
    sheets_note = []
    try:
        for sn in wb.sheetnames:
            action = sheet_action(sn)
            sheets_note.append(sn)
            if not action:
                continue
            ws = wb[sn]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [cell_str(c) for c in rows[0]]
            for r_i, row in enumerate(rows[1:], start=2):
                vals = list(row)
                if all(v is None or str(v).strip() == "" for v in vals):
                    continue
                while len(vals) < len(headers):
                    vals.append(None)
                row_map, _ = build_row_maps(headers, vals)

                project_code = row_map.get("项目代码", "")
                hilist_code = row_map.get("国家目录编码", "") or row_map.get("医疗目录编码", "")
                hilist_name = (
                    row_map.get("项目名称", "")
                    or row_map.get("医保目录名称", "")
                    or row_map.get("医疗服务项目名称", "")
                )
                begin_date = row_map.get("开始日期", "")
                end_date = row_map.get("终止日期", "")
                change_fields = normalize_change_fields(row_map.get("本期变更", ""))
                chrgitm_lv = map_chrgitm_lv(row_map.get("收费项目等级", ""))
                list_type = row_map.get("目录类别", "") or "201"

                if action == "ADD":
                    if not project_code:
                        continue
                    # 服务 ADD：仅 project_code，禁止 01* 进 hilist_code
                    out_hilist = ""
                    out_project = project_code
                else:
                    if not hilist_code:
                        continue
                    # 局发偶发把 01* 项目码填进「国家目录编码」列：按服务规则落 project_code
                    if hilist_code.startswith("01") and "-" not in hilist_code:
                        out_hilist = ""
                        out_project = hilist_code
                    else:
                        out_hilist = hilist_code
                        out_project = ""

                ext = {}
                for h, s in row_map.items():
                    if h in CORE_SKIP:
                        continue
                    if not useful_ext_value(s):
                        continue
                    # 尽量保留数值型限价
                    if h == "医保目录定价上限金额":
                        raw = None
                        for hh, vv in zip(headers, vals):
                            if hh == h and vv is not None:
                                raw = vv
                                break
                        if isinstance(raw, (int, float)):
                            ext[h] = int(raw) if float(raw) == int(raw) else raw
                        else:
                            try:
                                ext[h] = int(s) if "." not in s else float(s)
                            except ValueError:
                                ext[h] = s
                    elif h in (
                        "限制使用标志",
                        "医疗使用标志",
                        "生育使用标志",
                        "医保目录超限处理方式",
                        "辅助生殖标识",
                        "限住院使用标识",
                    ):
                        try:
                            ext[h] = int(s)
                        except ValueError:
                            ext[h] = s
                    else:
                        if h not in ext:
                            ext[h] = s

                lines.append(
                    {
                        "action": action,
                        "hilist_code": out_hilist,
                        "project_code": out_project,
                        "hilist_name": hilist_name,
                        "begin_date": begin_date,
                        "end_date": end_date,
                        "change_fields": change_fields,
                        "chrgitm_lv": chrgitm_lv,
                        "list_type": list_type,
                        "source_sheet": sn,
                        "source_row": str(r_i),
                        "ext_json": json.dumps(ext, ensure_ascii=False, separators=(",", ":")),
                    }
                )
    finally:
        wb.close()

    action_counts = Counter(x["action"] for x in lines)
    meta = {
        "doc_kind": "CATALOG_DELTA",
        "category": "SERVICE_F",
        "title": stem,
        "source_filename": SRC_NAME,
        "notice_date": "2026-07-28",
        "effective_date": "2026-07-31",
        "wash_note": (
            f"F服务局发目录更新（工程师版）；Sheet={sheets_note}；"
            "服务ADD仅写project_code(01*)；STOP/MODIFY用国家目录编码；"
            "STOP中5条国家目录编码实为01*项目码→改落project_code；"
            "ADD表有重复列名(价格构成/计价单位/计价说明)取首次非空；"
            "限价/超限等进ext_json；原件已在原始文件/未再move"
        ),
        "wash_tool": "cursor-ai",
    }

    (out_dir / "notice.meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    with open(out_dir / "notice.lines.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        for row in lines:
            w.writerow(row)

    summary = {
        "result": "OK",
        "std_dir": str(out_dir.relative_to(BASE)),
        "rows": len(lines),
        "actions": dict(action_counts),
        "doc_kind": "CATALOG_DELTA",
        "category": "SERVICE_F",
        "source_already_in_raw": True,
        "source": SRC_NAME,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    (BASE / "脚本" / "_wash_f_20260728_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
