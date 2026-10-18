# -*- coding: utf-8 -*-
"""Wash drug catalog xlsx into standard pack (notice.meta.json + notice.lines.csv)."""
from __future__ import annotations

import csv
import json
import os
import re
import shutil
import sys
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

BASE = Path(r"D:\0.product\02.Whale\whale-doc\whale-module-chs-jilin\目录更新通知")
OUT_ROOT = BASE / "标准化文件"
ARCHIVE = BASE / "原始文件"
FAIL_DIR = ARCHIVE / "_无法转换"

CSV_COLS = [
    "action",
    "hilist_code",
    "project_code",
    "hilist_name",
    "begin_date",
    "end_date",
    "change_fields",
    "chrgitm_lv",
    "list_type",
    "source_sheet",
    "source_row",
    "ext_json",
]

HILIST_KEYS = ("医疗目录编码", "药品代码", "目录编码", "国家目录编码", "中药配方颗粒代码")
NAME_KEYS = ("注册名称", "药品名称", "通用名", "项目名称", "医保目录名称")
BEGIN_KEYS = ("开始日期", "开始时间", "生效日期")
END_KEYS = ("结束日期", "终止日期", "停用日期")
LV_KEYS = ("收费项目等级",)
LIST_TYPE_KEYS = ("目录类别",)
CHANGE_KEYS = ("本期变更",)

CORE_CONSUMED = set(HILIST_KEYS + NAME_KEYS + BEGIN_KEYS + END_KEYS + LV_KEYS + LIST_TYPE_KEYS + CHANGE_KEYS)

ACTION_PREFIX = (
    ("新增", "ADD"),
    ("停用", "STOP"),
    ("省变更", "PROVINCE_MODIFY"),
    ("变更", "MODIFY"),
)


def parse_dates_from_filename(filename: str):
    dates = []
    # Y2024.10.17 / 2025.5.30 / 2026.3.6
    for m in re.finditer(r"(?:Y)?(\d{4})\.(\d{1,2})\.(\d{1,2})", filename):
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            dates.append(date(y, mo, d).isoformat())
        except ValueError:
            pass
    # lonely month.day after first date like 、3.16
    if dates:
        base_year = int(dates[0][:4])
        # match patterns like 、3.16 or 、04.2 that are not already captured as full yyyy.m.d
        for m in re.finditer(r"[、,，]\s*(\d{1,2})\.(\d{1,2})", filename):
            mo, d = int(m.group(1)), int(m.group(2))
            try:
                s = date(base_year, mo, d).isoformat()
                if s not in dates:
                    dates.append(s)
            except ValueError:
                pass
    notice_date = dates[0] if dates else None
    effective_date = dates[-1] if len(dates) > 1 else (dates[0] if dates else None)
    return notice_date, effective_date, dates


def sheet_action(sheet_name: str):
    for prefix, action in ACTION_PREFIX:
        if sheet_name.startswith(prefix):
            return action
    return None


def norm_cell(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def map_lv(v):
    if v is None:
        return ""
    s = str(v).strip()
    if not s or s.lower() == "none":
        return ""
    if s in ("1", "01", "甲"):
        return "1"
    if s in ("2", "02", "乙"):
        return "2"
    if s in ("3", "03", "丙", "自费"):
        return "3"
    # already 1/2/3
    if s in ("1", "2", "3"):
        return s
    return ""


def first_val(rowmap, keys):
    for k in keys:
        if k in rowmap and rowmap[k] not in (None, ""):
            return rowmap[k]
    return ""


def to_ext_json(rowmap: dict) -> str:
    ext = {}
    for k, v in rowmap.items():
        if k in CORE_CONSUMED:
            continue
        if v is None or v == "":
            continue
        if isinstance(v, (datetime, date)):
            ext[k] = norm_cell(v)
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            if isinstance(v, float) and v == int(v):
                ext[k] = int(v)
            else:
                ext[k] = v
        else:
            ext[k] = str(v).strip()
    return json.dumps(ext, ensure_ascii=False, separators=(",", ":"))


def title_from_filename(filename: str) -> str:
    name = Path(filename).stem
    name = re.sub(r"^附件[：:]\s*\d*\.?", "", name).strip()
    name = re.sub(r"^Y", "", name) if name.startswith("Y") and re.match(r"^Y\d", name) else name
    return name.strip() or filename


def wash_one(filename: str) -> dict:
    src = BASE / filename
    result = {
        "filename": filename,
        "status": "FAIL",
        "out_dir": "",
        "actions": {},
        "moved": False,
        "note": "",
    }
    if not src.exists():
        result["note"] = "根目录不存在"
        return result

    try:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    except Exception as e:
        FAIL_DIR.mkdir(parents=True, exist_ok=True)
        dest = FAIL_DIR / filename
        if dest.exists():
            ts = datetime.now().strftime("%Y%m%d%H%M%S")
            dest = FAIL_DIR / f"{src.stem}_{ts}{src.suffix}"
        shutil.move(str(src), str(dest))
        reason = FAIL_DIR / "reason.txt"
        with open(reason, "a", encoding="utf-8") as f:
            f.write(f"{filename}\t打不开: {e}\n")
        result["note"] = f"打不开: {e}"
        result["moved"] = True
        return result

    lines = []
    actions_count = {}
    known = False
    for sn in wb.sheetnames:
        action = sheet_action(sn)
        if not action:
            continue
        known = True
        ws = wb[sn]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            continue
        headers = [str(h).strip() if h is not None else "" for h in header]
        # skip blank trailing headers but keep index
        for row_idx, row in enumerate(it, start=2):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            rowmap = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                val = row[i] if i < len(row) else None
                rowmap[h] = norm_cell(val) if not isinstance(val, (int, float)) or isinstance(val, bool) else val
                # keep datetime normalized for date fields; for ext keep norm
                if isinstance(val, (datetime, date)):
                    rowmap[h] = norm_cell(val)
                elif val is None:
                    rowmap[h] = None
                elif isinstance(val, float) and val == int(val):
                    rowmap[h] = int(val)
                elif isinstance(val, (int, float)):
                    rowmap[h] = val
                else:
                    s = str(val).strip()
                    rowmap[h] = s if s else None

            hilist = first_val(rowmap, HILIST_KEYS)
            if not hilist:
                continue
            hilist = str(hilist).strip()
            name = first_val(rowmap, NAME_KEYS)
            begin = first_val(rowmap, BEGIN_KEYS)
            end = first_val(rowmap, END_KEYS)
            lv_raw = first_val(rowmap, LV_KEYS)
            lv = map_lv(lv_raw)
            list_type = first_val(rowmap, LIST_TYPE_KEYS)
            if list_type is not None and list_type != "":
                list_type = str(list_type).strip()
            else:
                list_type = "101"  # drug default
            change = first_val(rowmap, CHANGE_KEYS)
            change_fields = str(change).strip() if change else ""

            rec = {
                "action": action,
                "hilist_code": hilist,
                "project_code": "",
                "hilist_name": str(name) if name else "",
                "begin_date": str(begin) if begin else "",
                "end_date": str(end) if end else "",
                "change_fields": change_fields,
                "chrgitm_lv": lv,
                "list_type": list_type,
                "source_sheet": sn,
                "source_row": str(row_idx),
                "ext_json": to_ext_json(rowmap),
            }
            lines.append(rec)
            actions_count[action] = actions_count.get(action, 0) + 1

    wb.close()

    if not known or not lines:
        FAIL_DIR.mkdir(parents=True, exist_ok=True)
        dest = FAIL_DIR / filename
        if dest.exists():
            ts = datetime.now().strftime("%Y%m%d%H%M%S")
            dest = FAIL_DIR / f"{src.stem}_{ts}{src.suffix}"
        shutil.move(str(src), str(dest))
        with open(FAIL_DIR / "reason.txt", "a", encoding="utf-8") as f:
            f.write(f"{filename}\t无识别 Sheet 或无有效明细行\n")
        result["note"] = "无识别 Sheet 或无有效明细"
        result["moved"] = True
        return result

    # Windows strips trailing spaces in folder names; normalize stem
    stem = src.stem.rstrip(" .")
    out_dir = OUT_ROOT / stem
    out_dir.mkdir(parents=True, exist_ok=True)

    notice_date, effective_date, all_dates = parse_dates_from_filename(filename)
    meta = {
        "doc_kind": "CATALOG_DELTA",
        "category": "DRUG",
        "title": title_from_filename(filename),
        "source_filename": filename,
        "notice_date": notice_date,
        "effective_date": effective_date,
        "wash_note": (
            "药品局发目录更新；Sheet前缀映射ADD/STOP/MODIFY/PROVINCE_MODIFY；"
            "hilist_code取医疗目录编码/药品代码/目录编码；"
            "收费项目等级01/02/03或1/2/3或甲乙丙→1/2/3；"
            "限价与其余原列入ext_json；list_type缺省101。"
            + (f" 文件名多日期:{','.join(all_dates)}。" if len(all_dates) > 1 else "")
        ),
        "wash_tool": "cursor-ai",
    }
    with open(out_dir / "notice.meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
        f.write("\n")

    csv_path = out_dir / "notice.lines.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_COLS, quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        for rec in lines:
            w.writerow(rec)

    # archive original
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    dest = ARCHIVE / filename
    if dest.exists():
        ts = datetime.now().strftime("%Y%m%d%H%M%S")
        dest = ARCHIVE / f"{src.stem}_{ts}{src.suffix}"
    shutil.move(str(src), str(dest))

    result.update(
        {
            "status": "OK",
            "out_dir": str(out_dir),
            "actions": actions_count,
            "moved": True,
            "rows": len(lines),
            "archive": str(dest),
            "doc_kind": "CATALOG_DELTA",
            "category": "DRUG",
            "notice_date": notice_date,
        }
    )
    return result


def main():
    files = sys.argv[1:]
    if not files:
        print("usage: wash_drug_batch.py <file1> [file2...]")
        sys.exit(2)
    for fn in files:
        r = wash_one(fn)
        print(json.dumps(r, ensure_ascii=False))


if __name__ == "__main__":
    main()
