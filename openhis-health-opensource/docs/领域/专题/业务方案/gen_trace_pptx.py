# -*- coding: utf-8 -*-
"""生成住院药品追溯码方案院方会议 PPT"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.util import Inches, Pt, Emu
from pptx.oxml import parse_xml
from lxml import etree

OUT = Path(__file__).resolve().parent / "住院药品追溯码上传方案-院方会议.pptx"

# 视觉：医疗青绿 + 石板灰（避免紫/奶油风）
C = {
    "bg": "F4F7F6",
    "card": "FFFFFF",
    "ink": "1A2B2E",
    "muted": "5A6F74",
    "line": "D5E0DE",
    "teal": "0D7377",
    "teal_d": "085456",
    "teal_l": "E6F3F3",
    "accent": "C45C26",  # 陶土强调，少用
    "green": "2A9D6E",
    "amber": "C48A1A",
    "red": "B33A3A",
    "blue": "2B6A8F",
    "slate": "3D4F5F",
}


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def set_run(run, size=14, bold=False, color=C["ink"], font="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = font
    # 东亚字体
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if "ea" in child.tag or "latin" in child.tag:
            pass
    ea = parse_xml(
        f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
    )
    # avoid duplicate
    existing = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if existing is None:
        rPr.append(ea)


def add_text(shape, text, size=14, bold=False, color=C["ink"], align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return tf


def add_para(tf, text, size=13, bold=False, color=C["ink"], space_before=6):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return p


def fill_shape(shape, hex6):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(hex6)


def line_shape(shape, hex6=C["line"], pt=1):
    shape.line.color.rgb = rgb(hex6)
    shape.line.width = Pt(pt)


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def paint_bg(slide, color=C["bg"]):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    fill_shape(bg, color)
    line_shape(bg, color, 0)
    # send to back via z-order: recreate later shapes on top - ok for pptx


def header_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.95))
    fill_shape(bar, C["teal"])
    line_shape(bar, C["teal"], 0)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.95), Inches(13.333), Inches(0.06))
    fill_shape(accent, C["accent"])
    line_shape(accent, C["accent"], 0)
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12), Inches(0.45))
    add_text(t, title, size=26, bold=True, color="FFFFFF")
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.5), Inches(0.58), Inches(12), Inches(0.3))
        add_text(s, subtitle, size=12, color="D0E8E8")


def footer(slide, page, total=14):
    f = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(10), Inches(0.3))
    add_text(f, "住院药品追溯码采集与上传方案 · 院方讨论稿 · 内部资料", size=10, color=C["muted"])
    p = slide.shapes.add_textbox(Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3))
    add_text(p, f"{page} / {total}", size=10, color=C["muted"], align=PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, fill=C["card"]):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, fill)
    line_shape(sh, C["line"], 1)
    sh.adjustments[0] = 0.08
    return sh


def pill(slide, left, top, width, height, text, fill, font_size=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, fill)
    line_shape(sh, fill, 0)
    sh.adjustments[0] = 0.5
    add_text(sh, text, size=font_size, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    try:
        sh.text_frame.anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    return sh


def arrow_box(slide, left, top, width, height, text, fill=C["teal"]):
    sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, width, height)
    fill_shape(sh, fill)
    line_shape(sh, fill, 0)
    add_text(sh, text, size=11, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    try:
        sh.text_frame.anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    return sh


def flow_row(slide, y, items, colors=None):
    """横向流程箭头"""
    n = len(items)
    w = Inches(2.1)
    gap = Inches(0.15)
    total = n * w + (n - 1) * gap
    x0 = (Inches(13.333) - total) / 2
    colors = colors or [C["teal"]] * n
    for i, (txt, col) in enumerate(zip(items, colors)):
        arrow_box(slide, x0 + i * (w + gap), y, w, Inches(0.55), txt, col)


# ---------- slides ----------

def slide_cover(prs):
    slide = blank_slide(prs)
    # full teal left panel
    left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.2), Inches(7.5))
    fill_shape(left, C["teal"])
    line_shape(left, C["teal"], 0)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), 0, Inches(0.12), Inches(7.5))
    fill_shape(band, C["accent"])
    line_shape(band, C["accent"], 0)
    right = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.32), 0, Inches(8.01), Inches(7.5))
    fill_shape(right, C["bg"])
    line_shape(right, C["bg"], 0)

    t1 = slide.shapes.add_textbox(Inches(0.55), Inches(2.2), Inches(4.3), Inches(1.2))
    add_text(t1, "住院药品\n追溯码方案", size=36, bold=True, color="FFFFFF")
    t2 = slide.shapes.add_textbox(Inches(0.55), Inches(4.0), Inches(4.3), Inches(1))
    add_text(t2, "采集 · 绑定 · 退药清码\n出院统一上传医保", size=16, color="C8E6E6")

    r1 = slide.shapes.add_textbox(Inches(6.0), Inches(2.3), Inches(6.5), Inches(0.5))
    add_text(r1, "院方业务讨论稿", size=28, bold=True, color=C["ink"])
    r2 = slide.shapes.add_textbox(Inches(6.0), Inches(3.0), Inches(6.5), Inches(1.5))
    tf = add_text(r2, "供护理部 · 药剂科 · 医保办 · 信息科联合评审", size=14, color=C["muted"])
    add_para(tf, "涵盖五种落地路径：谁扫、何时扫、怎么扫、利弊与推荐组合", size=13, color=C["muted"])

    for i, label in enumerate(["方案1 护士摆药", "方案2 药房发药", "方案3 双扫", "方案4 瓶贴闭环", "方案5 分级组合"]):
        pill(slide, Inches(6.0), Inches(4.7) + Inches(i * 0.42), Inches(2.6), Inches(0.35), label, C["teal_d"] if i == 4 else C["slate"], 10)

    note = slide.shapes.add_textbox(Inches(9.0), Inches(6.6), Inches(3.8), Inches(0.5))
    add_text(note, "默认推荐：方案5 分级组合", size=12, bold=True, color=C["accent"])


def slide_agenda(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "会议议程", "Discussion Agenda")
    items = [
        ("01", "背景与目标", "监管要求、出院统一上传、拆零实物约束"),
        ("02", "五种方案速览", "谁扫 · 何时扫 · 一图看懂"),
        ("03", "分方案说明", "流程配图 + 利弊（方案1～5）"),
        ("04", "对比与推荐", "矩阵对比 · 方案5组合落地路径"),
        ("05", "需院方拍板事项", "目录、闸门、退药、打印门控"),
    ]
    for i, (no, title, desc) in enumerate(items):
        y = Inches(1.35) + Inches(i * 1.05)
        card(slide, Inches(0.6), y, Inches(12.1), Inches(0.9))
        num = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), y + Inches(0.18), Inches(0.55), Inches(0.55))
        fill_shape(num, C["teal"])
        line_shape(num, C["teal"], 0)
        add_text(num, no, size=14, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
        try:
            num.text_frame.anchor = MSO_ANCHOR.MIDDLE
        except Exception:
            pass
        tb = slide.shapes.add_textbox(Inches(1.7), y + Inches(0.18), Inches(10), Inches(0.35))
        add_text(tb, title, size=18, bold=True, color=C["ink"])
        db = slide.shapes.add_textbox(Inches(1.7), y + Inches(0.5), Inches(10), Inches(0.3))
        add_text(db, desc, size=13, color=C["muted"])
    footer(slide, page, total)


def slide_background(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "背景与共性约定", "为何要做、统一口径是什么")

    goals = [
        ("应扫尽扫", "监管品种发药链路须采集追溯码"),
        ("码实相符", "码与药品、批号、患者可关联"),
        ("退药闭环", "退多少清多少，避免幽灵码上传"),
        ("出院统传", "在院只存系统，结算时统一上传医保"),
    ]
    for i, (t, d) in enumerate(goals):
        x = Inches(0.5) + Inches(i * 3.15)
        c = card(slide, x, Inches(1.3), Inches(3.0), Inches(1.55), C["teal_l"])
        line_shape(c, C["teal"], 1)
        tb = slide.shapes.add_textbox(x + Inches(0.15), Inches(1.45), Inches(2.7), Inches(0.4))
        add_text(tb, t, size=16, bold=True, color=C["teal_d"])
        db = slide.shapes.add_textbox(x + Inches(0.15), Inches(1.95), Inches(2.7), Inches(0.7))
        add_text(db, d, size=12, color=C["ink"])

    # 实物铁律大图示
    card(slide, Inches(0.5), Inches(3.15), Inches(12.3), Inches(3.6))
    title = slide.shapes.add_textbox(Inches(0.75), Inches(3.3), Inches(11), Inches(0.35))
    add_text(title, "实物铁律：码在包装上，拆出的药没有码", size=16, bold=True, color=C["accent"])

    # mother box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.9), Inches(3.8), Inches(2.4))
    fill_shape(box, C["teal"])
    line_shape(box, C["teal_d"], 0)
    add_text(box, "", size=12)
    bt = slide.shapes.add_textbox(Inches(1.2), Inches(4.1), Inches(3.4), Inches(0.35))
    add_text(bt, "母盒 / 药瓶（唯一可扫）", size=13, bold=True, color="FFFFFF")
    code = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(4.6), Inches(3.0), Inches(0.7))
    fill_shape(code, "FFFFFF")
    line_shape(code, "FFFFFF", 0)
    add_text(code, "追溯码 ████████", size=14, bold=True, color=C["teal_d"], align=PP_ALIGN.CENTER)
    tip = slide.shapes.add_textbox(Inches(1.2), Inches(5.5), Inches(3.4), Inches(0.6))
    add_text(tip, "扫码永远对着包装\n不是对着药杯/药袋", size=12, color="E8FFFF")

    # arrow
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.1), Inches(4.8), Inches(1.0), Inches(0.45))
    fill_shape(arr, C["accent"])
    line_shape(arr, C["accent"], 0)

    # bags
    for i, name in enumerate(["李四药袋", "王五药袋", "赵六药袋"]):
        y = Inches(3.95) + Inches(i * 0.85)
        b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.4), y, Inches(2.8), Inches(0.7))
        fill_shape(b, C["card"])
        line_shape(b, C["line"], 1.5)
        add_text(b, f"{name} · 无追溯码", size=12, bold=True, color=C["muted"], align=PP_ALIGN.CENTER)

    note = slide.shapes.add_textbox(Inches(9.5), Inches(4.1), Inches(3.0), Inches(2.2))
    tf = add_text(note, "因此必须：", size=14, bold=True, color=C["ink"])
    add_para(tf, "① 先扫母盒，再分药", size=13)
    add_para(tf, "② 母盒用尽前不得扔", size=13)
    add_para(tf, "③ 退药扫母盒或带出原码", size=13)
    add_para(tf, "④ 禁止扫药袋「补码」", size=13, color=C["red"])

    footer(slide, page, total)


def slide_overview(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "五种方案一览", "Who · When · How 速览")

    rows = [
        ("方案1", "护士摆药扫", "药房不扫", "摆药时", "扫盒/母盒", C["blue"]),
        ("方案2", "药房发药扫", "护士不扫追溯", "发药时", "扫盒绑患者", C["slate"]),
        ("方案3", "药房+护士双扫", "两道核对", "核发+摆药", "池→归因", C["amber"]),
        ("方案4", "瓶贴闭环", "床旁不扫盒码", "打印前+床旁", "绑码→扫贴", C["green"]),
        ("方案5", "分级组合★", "默认推荐", "按品种", "2+4+3切片", C["teal"]),
    ]
    # header row
    headers = ["方案", "名称", "谁扫", "何时", "怎么扫（摘要）"]
    xs = [0.5, 2.0, 4.5, 7.2, 9.3]
    ws = [1.4, 2.4, 2.6, 2.0, 3.5]
    for x, w, h in zip(xs, ws, headers):
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.25), Inches(w), Inches(0.45))
        fill_shape(cell, C["teal_d"])
        line_shape(cell, C["teal_d"], 0)
        add_text(cell, h, size=12, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)

    for i, (a, b, c, d, e, col) in enumerate(rows):
        y = Inches(1.75) + Inches(i * 0.85)
        bgc = C["card"] if i % 2 == 0 else C["teal_l"]
        for x, w, txt in zip(xs, ws, [a, b, c, d, e]):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), y, Inches(w), Inches(0.75))
            fill_shape(cell, col if x == xs[0] else bgc)
            line_shape(cell, C["line"], 0.75)
            add_text(cell, txt, size=12, bold=(x == xs[0]), color="FFFFFF" if x == xs[0] else C["ink"], align=PP_ALIGN.CENTER)

    tip = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.6))
    add_text(tip, "★ 方案5 = 整盒走方案2 + 拆零走方案4 + 高值走方案3，多数医院会议建议以此为讨论基线。", size=13, bold=True, color=C["teal_d"])
    footer(slide, page, total)


def scheme_slide(prs, page, total, no, name, who, when, how_title, steps, pros, cons, color, flow):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, f"{no}  {name}", f"谁扫：{who} ｜ 何时：{when}")

    # flow
    flow_row(slide, Inches(1.25), flow, [color] * len(flow))

    # how card
    card(slide, Inches(0.45), Inches(2.05), Inches(7.6), Inches(4.7))
    t = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(7), Inches(0.35))
    add_text(t, how_title, size=15, bold=True, color=color)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(2.65), Inches(7.1), Inches(3.9))
    tf = add_text(body, steps[0], size=13, color=C["ink"])
    for s in steps[1:]:
        add_para(tf, s, size=13, space_before=8)

    # pros
    card(slide, Inches(8.25), Inches(2.05), Inches(4.6), Inches(2.2), C["teal_l"])
    pt = slide.shapes.add_textbox(Inches(8.45), Inches(2.15), Inches(4.2), Inches(0.3))
    add_text(pt, "利", size=14, bold=True, color=C["green"])
    pb = slide.shapes.add_textbox(Inches(8.45), Inches(2.5), Inches(4.2), Inches(1.6))
    tf = add_text(pb, "· " + pros[0], size=12)
    for p in pros[1:]:
        add_para(tf, "· " + p, size=12, space_before=4)

    # cons
    card(slide, Inches(8.25), Inches(4.45), Inches(4.6), Inches(2.3), "FBF0EC")
    ct = slide.shapes.add_textbox(Inches(8.45), Inches(4.55), Inches(4.2), Inches(0.3))
    add_text(ct, "弊", size=14, bold=True, color=C["red"])
    cb = slide.shapes.add_textbox(Inches(8.45), Inches(4.9), Inches(4.2), Inches(1.7))
    tf = add_text(cb, "· " + cons[0], size=12)
    for p in cons[1:]:
        add_para(tf, "· " + p, size=12, space_before=4)

    footer(slide, page, total)


def slide_scheme5_diagram(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "方案5 分级组合 · 配图", "一张图看清：什么药走哪条扫码路径")

    # tree root
    root = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.25), Inches(4.3), Inches(0.65))
    fill_shape(root, C["teal"])
    line_shape(root, C["teal"], 0)
    add_text(root, "监管品种？", size=16, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)

    # branches
    cards_data = [
        (0.4, 2.3, C["slate"], "非监管", "免扫 / 抽查\n（院方配置）", "卡片 D"),
        (3.5, 2.3, C["blue"], "监管 · 整盒", "药房发药扫\n码绑患者", "卡片 A = 方案2"),
        (6.6, 2.3, C["green"], "监管 · 拆零", "打印前扫母盒\n床旁扫瓶贴", "卡片 B = 方案4"),
        (9.7, 2.3, C["amber"], "高值 / 精麻", "药房入池\n护士再归因", "卡片 C = 方案3"),
    ]
    for x, y, col, title, body, tag in cards_data:
        c = card(slide, Inches(x), Inches(y), Inches(2.9), Inches(2.5))
        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2.9), Inches(0.45))
        fill_shape(top, col)
        line_shape(top, col, 0)
        add_text(top, title, size=13, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
        b = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.65), Inches(2.6), Inches(1.2))
        add_text(b, body, size=13, color=C["ink"], align=PP_ALIGN.CENTER)
        tg = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 1.95), Inches(2.6), Inches(0.35))
        add_text(tg, tag, size=11, bold=True, color=col, align=PP_ALIGN.CENTER)

    # bottom timeline
    card(slide, Inches(0.5), Inches(5.1), Inches(12.3), Inches(1.6))
    tt = slide.shapes.add_textbox(Inches(0.7), Inches(5.2), Inches(11), Inches(0.3))
    add_text(tt, "推荐落地节奏（降低排班冲击）", size=14, bold=True, color=C["ink"])
    phases = [
        ("第一阶段", "监管整盒上方案2", C["blue"]),
        ("第二阶段", "拆零上方案4瓶贴闭环", C["green"]),
        ("第三阶段", "高值加方案3双扫", C["amber"]),
    ]
    for i, (a, b, col) in enumerate(phases):
        x = Inches(0.9) + Inches(i * 4.0)
        pill(slide, x, Inches(5.7), Inches(1.5), Inches(0.35), a, col, 11)
        tb = slide.shapes.add_textbox(x + Inches(1.6), Inches(5.7), Inches(2.2), Inches(0.35))
        add_text(tb, b, size=12, color=C["ink"])
    tip = slide.shapes.add_textbox(Inches(0.7), Inches(6.25), Inches(11.5), Inches(0.3))
    add_text(tip, "HIS 发药/摆药界面建议直接提示本行策略（A/B/C），避免口头记忆出错。", size=12, color=C["muted"])
    footer(slide, page, total)


def slide_compare(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "综合对比矩阵", "从工作量、合规、拆零友好度看差异")

    dims = ["护士工作量", "药房工作量", "拆零无码袋", "监管强度", "设备投入", "落地速度"]
    # scores as text stars / labels
    data = {
        "方案1": ["高", "低", "高", "中", "低", "快"],
        "方案2": ["低", "高", "中", "中高", "低", "快"],
        "方案3": ["高", "高", "最高", "最高", "中", "中"],
        "方案4": ["中", "中", "高", "高", "高", "中"],
        "方案5★": ["中（削峰）", "中（削峰）", "高", "高可配", "中高", "可分期"],
    }
    schemes = list(data.keys())
    # corner
    cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.25), Inches(2.0), Inches(0.5))
    fill_shape(cell, C["teal_d"])
    line_shape(cell, C["teal_d"], 0)
    add_text(cell, "维度 \\ 方案", size=11, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)

    for j, s in enumerate(schemes):
        x = Inches(2.4) + Inches(j * 2.1)
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.25), Inches(2.1), Inches(0.5))
        fill_shape(cell, C["teal"] if "5" in s else C["slate"])
        line_shape(cell, C["line"], 0)
        add_text(cell, s, size=12, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)

    for i, d in enumerate(dims):
        y = Inches(1.75) + Inches(i * 0.75)
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), y, Inches(2.0), Inches(0.75))
        fill_shape(cell, C["teal_l"])
        line_shape(cell, C["line"], 0.75)
        add_text(cell, d, size=11, bold=True, color=C["ink"], align=PP_ALIGN.CENTER)
        for j, s in enumerate(schemes):
            x = Inches(2.4) + Inches(j * 2.1)
            bg = C["card"] if i % 2 == 0 else "EEF5F5"
            if "5" in s:
                bg = "E6F3F3"
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.1), Inches(0.75))
            fill_shape(cell, bg)
            line_shape(cell, C["line"], 0.75)
            add_text(cell, data[s][i], size=12, color=C["ink"], align=PP_ALIGN.CENTER)

    footer(slide, page, total)


def slide_recommend(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "会议建议结论", "给院领导的一页纸")

    big = card(slide, Inches(0.5), Inches(1.3), Inches(12.3), Inches(1.4), C["teal"])
    line_shape(big, C["teal"], 0)
    t = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(1.0))
    tf = add_text(t, "建议以「方案5 分级组合」为院内基线方案", size=22, bold=True, color="FFFFFF")
    add_para(tf, "整盒 → 方案2　　拆零 → 方案4　　高值/精麻 → 方案3　　非监管 → 可免扫", size=14, color="D0EEEE", space_before=10)

    cols = [
        ("大多数医院", "方案5\n（2+4+3）", C["teal"]),
        ("最快合规", "先上方案2\n整盒切片", C["blue"]),
        ("已有PDA+瓶贴", "优先打满\n方案4", C["green"]),
        ("飞检极严", "高值加\n方案3双扫", C["amber"]),
    ]
    for i, (a, b, col) in enumerate(cols):
        x = Inches(0.5) + Inches(i * 3.2)
        c = card(slide, x, Inches(3.0), Inches(3.0), Inches(2.4))
        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(3.0), Inches(3.0), Inches(0.5))
        fill_shape(top, col)
        line_shape(top, col, 0)
        add_text(top, a, size=13, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
        body = slide.shapes.add_textbox(x + Inches(0.2), Inches(3.7), Inches(2.6), Inches(1.5))
        add_text(body, b, size=16, bold=True, color=C["ink"], align=PP_ALIGN.CENTER)

    note = slide.shapes.add_textbox(Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.0))
    tf = add_text(note, "不推荐：全院强制床旁扫盒码（拆零必败）；只扫不绑患者；发药即传医保替代出院统传。", size=13, color=C["red"])
    add_para(tf, "上传时点建议维持：在院期间码存 HIS，出院结算统一上传——便于退药改码与共用码处理。", size=13, color=C["muted"], space_before=8)
    footer(slide, page, total)


def slide_decisions(prs, page, total):
    slide = blank_slide(prs)
    paint_bg(slide)
    header_bar(slide, "请院方拍板", "会议决议清单（可勾选）")

    items = [
        "主路径是否采用方案5分级组合（推荐）？若否，单选哪一方案？",
        "监管整盒：是否强制「发 N 盒 = 扫 N 码」？",
        "拆零：是否强制「先扫母盒再分袋/打贴」？母盒架是否设立？",
        "是否上瓶贴/医嘱单打印门控（无追溯码禁止打印）？",
        "退药：有母盒再扫；仅裸药是否允许「带出原码」+双签？",
        "出院「应扫未扫」：提醒 还是 阻断结算？",
        "多人共用一码的医保上传计数规则（医保办）？",
        "抢救/手术室免扫与补录时限（如 24 小时）？",
    ]
    for i, text in enumerate(items):
        y = Inches(1.25) + Inches(i * 0.65)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), y, Inches(0.4), Inches(0.4))
        fill_shape(box, C["card"])
        line_shape(box, C["teal"], 1.5)
        tb = slide.shapes.add_textbox(Inches(1.15), y + Inches(0.05), Inches(11.5), Inches(0.4))
        add_text(tb, f"{i+1}. {text}", size=13, color=C["ink"])
    footer(slide, page, total)


def slide_end(prs):
    slide = blank_slide(prs)
    left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    fill_shape(left, C["teal"])
    line_shape(left, C["teal"], 0)
    t = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10), Inches(1))
    add_text(t, "谢谢 · 请讨论", size=40, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    s = slide.shapes.add_textbox(Inches(1.5), Inches(3.7), Inches(10), Inches(1.2))
    tf = add_text(s, "文档位置：whale-doc/health/业务方案/", size=16, color="C8E6E6", align=PP_ALIGN.CENTER)
    add_para(tf, "总册 + 方案1～5分册（谁扫/何时扫/怎么扫/利弊）", size=14, color="C8E6E6", space_before=8)
    add_para(tf, "默认讨论基线：方案5 = 整盒② + 拆零④ + 高值③", size=14, color="FFFFFF", space_before=12)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 14

    slide_cover(prs)
    slide_agenda(prs, 2, total)
    slide_background(prs, 3, total)
    slide_overview(prs, 4, total)

    scheme_slide(
        prs, 5, total,
        "方案1", "护士摆药扫码（药房不扫）",
        "摆药护士主责", "摆药分到患者前",
        "怎么扫 · 正向与逆向",
        [
            "整盒：连扫 N 个盒码 → 绑患者（码数=盒数）",
            "拆零：先扫母盒 → 多人可共用同一码 → 再分入无码药袋 → 母盒归架",
            "退药：再扫被退盒/母盒；仅裸药则系统带出原码确认",
            "禁止：扫药袋补码；一码顶多盒；退药不清码",
            "口诀：药房发药不扫码；摆药台上扫盒码；先扫后分母盒留",
        ],
        ["药房零增量，发药不受拖累", "拆零归因直观，病区退药可闭环", "设备投入低，落地快"],
        ["护士高峰工作量大，易漏扫", "无药房码池，飞检说服力偏弱", "依赖母盒架，乱扔包装难再扫"],
        C["blue"],
        ["药房发药", "摆药扫盒码", "绑到患者", "出院统传"],
    )

    scheme_slide(
        prs, 6, total,
        "方案2", "药房发药扫码（护士不扫追溯）",
        "药房药师主责", "发药/核发当时",
        "怎么扫 · 正向与逆向",
        [
            "整盒：窗口连扫 N 码 → 写入该患者发放行",
            "拆零：扫母盒码 → 必须分别写到每个患者行（允许共用）",
            "护士：不扫追溯码，但须保管母盒、整盒退回原盒",
            "退药：原盒/母盒回药房再扫清码；无盒走带出原码",
            "口诀：发药窗口扫盒码；码要落到每个人；退药原盒回药房",
        ],
        ["护理零增量，落地阻力小", "出库即有码，合规口径清晰", "整盒窗口连扫专业高效"],
        ["药房成瓶颈，高峰易排队", "拆零归因是难点（必须到人）", "退药强依赖原盒回药房"],
        C["slate"],
        ["药房扫码", "绑到患者", "病区用药", "退药回扫", "出院统传"],
    )

    scheme_slide(
        prs, 7, total,
        "方案3", "药房 + 护士双扫",
        "药师第一道 + 护士第二道", "核发入池 → 摆药归因",
        "怎么扫 · 两道分工",
        [
            "药房：扫追溯码进入汇总码池（回答：发出哪些码？）",
            "护士：仅扫/选池内码归因到患者（回答：码给了谁？）",
            "拆零：药房扫母盒入池 → 护士持母盒确认后多人绑定 → 再分袋",
            "退药：护士发起带码 + 药房确认入库（建议终态在药房）",
            "口诀：药房先扫入码池；护士再扫落到人；池外码一律拒",
        ],
        ["双重核对，飞检说服力最强", "防串码，职责分离清晰", "与现有码池+护士分配产品贴近"],
        ["人力成本最高，流程最长", "两端积压会卡住全链路", "全品种双扫过重，宜限高值"],
        C["amber"],
        ["药房入池", "护士归因", "绑到患者", "双侧退药", "出院统传"],
    )

    scheme_slide(
        prs, 8, total,
        "方案4", "瓶贴 / 医嘱单闭环（推荐4D）",
        "打印前绑追溯码；床旁扫业务码", "打印前 + 床旁给药",
        "怎么扫 · 两种码不要混淆",
        [
            "① 摆药/打印前：扫母盒国家追溯码 → 绑患者（监管无码则拒打）",
            "② 打印瓶贴或医嘱执行单（业务条码关联发放行）",
            "③ 床旁：扫腕带 + 瓶贴/医嘱单 —— 不扫盒上追溯码",
            "④ 拆零：一人一贴，库内可共用同一母盒追溯码；母盒不必随车",
            "口诀：打印之前扫盒码；有码才打瓶贴单；床旁只扫腕带贴",
        ],
        ["完美避开拆零药袋无码", "闭环护理体验好，易推广", "整盒/拆零床旁动作统一"],
        ["需 PDA+瓶贴机，改造成本高", "错贴/漏打仍有给错药风险", "退药不能只撕贴，仍要清追溯码"],
        C["green"],
        ["扫母盒绑码", "打印瓶贴", "床旁扫贴", "退药清码", "出院统传"],
    )

    slide_scheme5_diagram(prs, 9, total)

    # scheme5利弊简页
    scheme_slide(
        prs, 10, total,
        "方案5", "分级组合（默认推荐）",
        "按品种分别套用 2 / 4 / 3", "见策略卡片 A/B/C",
        "为什么比「四选一」更好",
        [
            "整盒走方案2：不拖垮护理，窗口快扫合规",
            "拆零走方案4：不逼床旁找盒码，瓶贴闭环",
            "高值走方案3：人力打在风险最高的品种上",
            "可分期：先 A 再 B 再 C，排班冲击可控",
            "代价：制度/系统要支持按品种策略，培训分场景",
        ],
        ["多数医院最优平衡点", "飞检口径清晰可配置", "继承各积木优点、削峰填谷"],
        ["SOP 不能一张打天下", "系统改造大于单方案", "初期易混用规则（需界面提示）"],
        C["teal"],
        ["判品种", "走对应卡片", "分场景扫", "统一出院传"],
    )

    slide_compare(prs, 11, total)
    slide_recommend(prs, 12, total)
    slide_decisions(prs, 13, total)
    slide_end(prs)

    prs.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
